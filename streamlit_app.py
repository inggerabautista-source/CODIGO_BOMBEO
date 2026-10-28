# -*- coding: utf-8 -*-
"""
MEMORIA DESCRIPTIVA Y DE CÁLCULO DE LÍNEA DE BOMBEO
Criterios basados en los manuales MAPAS - CONAGUA (Datos Básicos, Conducción,
Líneas de conducción, Pozos, Cárcamos, Tanques de regulación).

Ejecutar:   pip install streamlit pandas numpy matplotlib openpyxl python-docx
            streamlit run memoria_bombeo.py

IMPORTANTE: todas las tablas precargadas (rugosidades, K, coeficientes de regulación,
tiempos de ciclo, etc.) son VALORES DE REFERENCIA EDITABLES. Verifica cada uno
contra la edición vigente del manual MAPAS antes de emitir una memoria oficial.
"""
import csv
import io
import math
import re
import textwrap
import datetime as dt

import numpy as np
import pandas as pd
import streamlit as st
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

G = 9.81          # m/s2
NU = 1.01e-6      # viscosidad cinemática agua 20 °C (m2/s)

# ----------------------------------------------------------------------------
# TABLAS PRECARGADAS (editar/verificar con MAPAS)
# eps: rugosidad absoluta (mm) | C: Hazen-Williams | f_di: DI/DN aprox. | a: celeridad (m/s)
# ----------------------------------------------------------------------------
MATERIALES = {
    "PVC hidráulico":            dict(eps=0.0015, C=150, f_di=0.95, a=400),
    "PEAD (polietileno AD)":     dict(eps=0.0015, C=150, f_di=0.90, a=300),
    "Acero soldado (nuevo)":     dict(eps=0.10,   C=120, f_di=0.98, a=1000),
    "Acero galvanizado":         dict(eps=0.15,   C=120, f_di=1.00, a=1000),
    "Fierro dúctil c/ mortero":  dict(eps=0.12,   C=130, f_di=1.00, a=1100),
    "Fierro fundido (nuevo)":    dict(eps=0.25,   C=130, f_di=1.00, a=1100),
    "Fibrocemento (A-C)":        dict(eps=0.025,  C=135, f_di=1.00, a=900),
    "Concreto presforzado":      dict(eps=0.30,   C=120, f_di=1.00, a=1000),
    "PRFV (fibra de vidrio)":    dict(eps=0.01,   C=140, f_di=0.98, a=450),
}
DN_COMERCIALES = [50, 63, 75, 100, 150, 200, 250, 300, 350, 400, 450, 500, 600,
                  750, 900, 1050, 1200, 1500]

K_TABLE = {
    "Entrada a tubería (borde vivo)": 0.50,
    "Codo 90°": 0.90,
    "Codo 90° radio largo": 0.60,
    "Codo 45°": 0.40,
    "Codo 22.5°": 0.20,
    "Tee paso directo": 0.60,
    "Tee salida lateral": 1.80,
    "Válvula de compuerta (abierta)": 0.20,
    "Válvula de mariposa (abierta)": 0.50,
    "Válvula de retención (columpio)": 2.50,
    "Válvula de retención (doble disco)": 1.50,
    "Válvula de pie con coladera": 1.75,
    "Reducción gradual": 0.15,
    "Ampliación gradual": 0.30,
    "Medidor electromagnético": 0.00,
    "Junta flexible / Dresser": 0.00,
    "Válvula de expulsión de aire": 0.00,
    "Válvula de desagüe (purga)": 0.00,
    "Salida a tanque (descarga)": 1.00,
}
ACC_OPTIONS = list(K_TABLE) + ["Otro (manual)"]
HP_COMERCIALES = [1, 1.5, 2, 3, 5, 7.5, 10, 15, 20, 25, 30, 40, 50, 60, 75, 100,
                  125, 150, 200, 250, 300, 350, 400, 500]
PN_KGCM2 = [5, 7, 10, 14, 20, 25, 32, 40]

# Coeficiente de regulación (m3 por l/s de Qmd). Valor de referencia 24 h = 14.58
# (MAPAS - Datos Básicos). Para otros horarios usa la curva masa (pestaña Tanques).
CR_24H = 14.58
PATRON_DEMANDA = [0.45, 0.45, 0.45, 0.45, 0.45, 0.60, 0.90, 1.35, 1.50, 1.50, 1.40, 1.40,
                  1.40, 1.40, 1.40, 1.35, 1.25, 1.20, 1.20, 1.10, 0.90, 0.70, 0.55, 0.50]
# ^ PATRÓN GENÉRICO de ejemplo; sustituir por el de la ley de demanda del manual / local.

DEST_ESTR = {"Cisterna + tanque de piso": ["Cisterna", "Tanque de piso"],
             "Tanque de piso + tanque elevado": ["Tanque de piso", "Tanque elevado"],
             "Cisterna + tanque de piso + tanque elevado": ["Cisterna", "Tanque de piso", "Tanque elevado"]}
DEST_NUEVOS = list(DEST_ESTR)

EQUIPOS = {
    "Pozo profundo": ["Sumergible", "Turbina vertical (flecha larga)"],
    "Cárcamo de bombeo": ["Sumergible (cárcamo húmedo)", "Centrífuga horizontal (cárcamo seco)",
                          "Turbina vertical (flecha larga)"],
}
ETA_B = {"Sumergible": 0.68, "Turbina vertical (flecha larga)": 0.72,
         "Sumergible (cárcamo húmedo)": 0.68, "Centrífuga horizontal (cárcamo seco)": 0.70}


# ----------------------------------------------------------------------------
# FUNCIONES HIDRÁULICAS
# ----------------------------------------------------------------------------
def hf_pipe(Q, D, L, eps_mm, C, method):
    """Pérdida por fricción en un tramo. Q m3/s, D m, L m."""
    if L <= 0 or Q <= 0:
        return 0.0
    if method.startswith("Hazen"):
        return 10.67 * L * Q ** 1.852 / (C ** 1.852 * D ** 4.87)
    v = 4 * Q / (math.pi * D ** 2)
    Re = max(v * D / NU, 1.0)
    e = eps_mm / 1000.0
    f = 0.25 / (math.log10(e / (3.7 * D) + 5.74 / Re ** 0.9)) ** 2   # Swamee-Jain
    if Re < 2300:
        f = 64 / Re
    return f * (L / D) * v ** 2 / (2 * G)


def commercial_up(x_mm, lst=DN_COMERCIALES):
    for d in lst:
        if d >= x_mm - 1e-9:
            return d
    return lst[-1]


def parse_ch(v):
    """Cadenamiento a metros: acepta 1250, '1250,5' o '1+250'. NaN si no es válido."""
    try:
        if isinstance(v, str):
            t = v.strip().replace(" ", "")
            if "+" in t:
                km, m = t.split("+")
                return float(km.replace(",", "")) * 1000 + float(m.replace(",", "."))
            return float(t.replace(",", "."))
        return float(v)
    except (ValueError, TypeError):
        return float("nan")


def _num(v):
    """Número desde texto con coma decimal o separador de miles; NaN si no es numérico."""
    if v is None:
        return float("nan")
    if isinstance(v, (int, float, np.number)):
        return float(v)
    t = str(v).strip().replace(" ", "").replace("\u00a0", "")
    if not t:
        return float("nan")
    if "," in t and "." in t:
        t = t.replace(",", "")           # 2,240.50
    elif "," in t:
        t = t.replace(",", ".")          # 2240,50
    try:
        return float(t)
    except ValueError:
        return float("nan")


def read_table(file):
    """Lee CSV (separador y codificación automáticos) o Excel, ubicando la fila de encabezado
    aunque haya títulos arriba. Todo se lee como texto; la conversión numérica se hace después."""
    raw = file.getvalue() if hasattr(file, "getvalue") else file.read()
    if file.name.lower().endswith(".csv"):
        try:
            txt = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            txt = raw.decode("latin-1")
        lines = [ln for ln in txt.splitlines() if ln.strip()]
        sample = lines[:40]
        delim = ","
        for d in (";", "\t", "|"):
            if sample and sum(d in ln for ln in sample) >= len(sample):
                delim = d
                break
        rows = list(csv.reader(lines, delimiter=delim))
        w = max(len(r) for r in rows)
        raw_df = pd.DataFrame([r + [""] * (w - len(r)) for r in rows])
    else:
        raw_df = pd.read_excel(io.BytesIO(raw), header=None, dtype=str).fillna("")
    raw_df = raw_df.astype(str)
    raw_df = raw_df[~(raw_df.apply(lambda c: c.str.strip()) == "").all(axis=1)].reset_index(drop=True)
    h, hallado = 0, False
    for i in range(min(len(raw_df), 25)):
        vals = [v.lower() for v in raw_df.iloc[i].tolist()]
        if any(w_ in v for v in vals for w_ in ("cad", "cota", "elev", "km", "dist", "estac")):
            h, hallado = i, True
            break
    if not hallado:
        # Sin encabezado: si la primera fila ya son números (p. ej. 0+023.20 | 283.58), es un dato, no un título
        n_num = sum(1 for v in raw_df.iloc[0].tolist() if not np.isnan(parse_ch(v)))
        if n_num >= 2:
            df = raw_df.copy()
            df.columns = [f"columna {j + 1}" for j in range(df.shape[1])]
            return df
    names, seen = [], {}
    for j, v in enumerate(raw_df.iloc[h].tolist()):
        v = v.strip() or f"columna {j + 1}"
        seen[v] = seen.get(v, 0) + 1
        names.append(v if seen[v] == 1 else f"{v} ({seen[v]})")
    df = raw_df.iloc[h + 1:].reset_index(drop=True)
    df.columns = names
    return df


def read_profile(file, kind="terreno", key="prf"):
    """Perfil (cadenamiento, cota). Detecta columnas, deja elegirlas y omite filas no numéricas.
    Devuelve DataFrame (ch, z) o None si no hay datos válidos."""
    try:
        df = read_table(file)
    except Exception as ex:
        st.error(f"No se pudo leer el archivo: {ex}")
        return None
    cols = list(df.columns)
    low = {c: str(c).lower() for c in cols}

    def pick(words, exclude=()):
        for c, l in low.items():
            if any(w in l for w in words) and not any(x in l for x in exclude):
                return c
        return None

    c_ch = pick(["cad", "km", "dist", "estac"]) or cols[0]
    if kind == "tuberia":
        c_z = (pick(["tub", "rasante", "proyecto", "clave", "lomo", "eje", "plantilla"])
               or pick(["cota", "elev"], exclude=["terreno", "natural"]))
    else:
        c_z = pick(["terreno", "natural"]) or pick(["cota", "elev"]) or pick(["z"])
    if c_z is None or c_z == c_ch:
        c_z = next((c for c in cols if c != c_ch), cols[0])
    k1, k2 = st.columns(2)
    c_ch = k1.selectbox(f"Columna de cadenamiento ({kind})", cols, index=cols.index(c_ch), key=f"{key}_ch")
    c_z = k2.selectbox(f"Columna de cota ({kind})", cols, index=cols.index(c_z), key=f"{key}_z")
    out = pd.DataFrame({"ch": df[c_ch].map(parse_ch), "z": df[c_z].map(_num)})
    malas = int(out.isna().any(axis=1).sum())
    out = out.dropna().sort_values("ch").reset_index(drop=True)
    if malas:
        st.warning(f"Se omitieron {malas} fila(s) sin valor numérico en «{c_ch}» o «{c_z}» "
                   "(encabezados repetidos, unidades, celdas vacías o texto).")
    if len(out) < 2:
        st.error(f"Con las columnas «{c_ch}» y «{c_z}» no hay al menos 2 filas numéricas. "
                 "Elige otras columnas o revisa el archivo.")
        return None
    return out


def demo_profile():
    ch = np.arange(0, 2001, 100.0)
    z = 2240 + 0.022 * ch + 9 * np.sin(ch / 260.0) + 6 * np.sin(ch / 90.0)
    return pd.DataFrame({"ch": ch, "z": np.round(z, 2)})


def hidraulica(prof, cover, D, Q, eps, C, method, acc, z_dest):
    ch = prof["ch"].to_numpy(float)
    zt = prof["z"].to_numpy(float)
    zp = prof["zp"].to_numpy(float) if "zp" in prof else zt - cover
    Lseg = np.hypot(np.diff(ch), np.diff(zp))
    hfs = np.array([hf_pipe(Q, D, l, eps, C, method) for l in Lseg])
    a = acc.copy()
    a = a.dropna(subset=["Cadenamiento (m)", "Elemento"]) if len(a) else a
    rows = []
    for _, r in a.iterrows():
        K = r["K"] if pd.notna(r["K"]) else K_TABLE.get(r["Elemento"], 0.0)
        n = r["Cantidad"] if pd.notna(r["Cantidad"]) else 1
        d = (r["Diámetro (mm)"] / 1000.0) if pd.notna(r["Diámetro (mm)"]) and r["Diámetro (mm)"] > 0 else D
        v = Q / (math.pi * d ** 2 / 4)
        rows.append(dict(c=float(r["Cadenamiento (m)"]), Elemento=r["Elemento"], n=n, K=K,
                         d_mm=d * 1000, v=v, hl=n * K * v ** 2 / (2 * G)))
    accr = pd.DataFrame(rows, columns=["c", "Elemento", "n", "K", "d_mm", "v", "hl"])
    down_hf = np.array([hfs[i:].sum() for i in range(len(ch))])
    down_hl = np.array([accr.loc[accr["c"] >= c - 1e-6, "hl"].sum() for c in ch])
    LE = z_dest + down_hf + down_hl
    return dict(ch=ch, zt=zt, zp=zp, Lseg=Lseg, hfs=hfs, accr=accr, LE=LE,
                P=LE - zp, Pst=z_dest - zp, hf=hfs.sum(), hl=accr["hl"].sum(),
                H0=LE[0], L=Lseg.sum(), v=Q / (math.pi * D ** 2 / 4))


def puntos_alto_bajo(zp):
    hi, lo = [], []
    for i in range(1, len(zp) - 1):
        if zp[i] > zp[i - 1] and zp[i] >= zp[i + 1]:
            hi.append(i)
        if zp[i] < zp[i - 1] and zp[i] <= zp[i + 1]:
            lo.append(i)
    return hi, lo


def potencia(Q, H, eta_b, eta_m, fs=1.15):
    ph = G * Q * H                              # kW (Q m3/s * 9.81 * H) = kW
    p_eje = ph / eta_b
    p_motor = p_eje * fs
    hp_req = p_motor / 0.7457
    hp_com = next((h for h in HP_COMERCIALES if h >= hp_req), HP_COMERCIALES[-1])
    return dict(ph=ph, p_eje=p_eje, p_elec=p_eje / eta_m, hp_req=hp_req, hp_com=hp_com)


def cr_curva_masa(patron, T, h0):
    p = np.array(patron, float)
    p = p / p.sum()
    entra = np.zeros(24)
    for k in range(int(round(T))):
        entra[(h0 + k) % 24] = 1.0 / T
    cum = np.cumsum(entra - p)
    frac = cum.max() - cum.min()
    return frac * 86.4, cum        # m3 por l/s de Qmd


# ----------------------------------------------------------------------------
# GRÁFICAS
# ----------------------------------------------------------------------------
def fig_perfil(prof, res, acc_df, hi, lo, pn_m=None):
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(11, 7), sharex=True,
                                 gridspec_kw={"height_ratios": [3, 1.3]})
    ch = res["ch"]
    a1.plot(ch, res["zt"], color="saddlebrown", lw=1.5, label="Terreno")
    a1.plot(ch, res["zp"], color="navy", lw=1.8, label="Tubería")
    a1.plot(ch, res["LE"], color="red", ls="--", lw=1.5, label="Línea de energía (bombeo)")
    a1.axhline(res["LE"][-1] - res["accr"]["hl"].tail(1).sum() if len(res["accr"]) else res["LE"][-1],
               color="c", lw=0.6, ls=":")
    a1.fill_between(ch, res["zt"], res["zt"].min() - 5, color="burlywood", alpha=0.3)
    if len(hi):
        a1.scatter(ch[hi], res["zp"][hi], marker="^", color="orange", s=70, zorder=5, label="Puntos altos")
    if len(lo):
        a1.scatter(ch[lo], res["zp"][lo], marker="v", color="green", s=70, zorder=5, label="Puntos bajos")
    markers = {"Válvula de expulsión de aire": ("*", "darkorange", 160),
               "Válvula de desagüe (purga)": ("s", "green", 60),
               "Válvula de retención (columpio)": ("D", "purple", 50),
               "Válvula de compuerta (abierta)": ("X", "black", 55)}
    seen = set()
    for _, r in res["accr"].iterrows():
        zc = np.interp(r["c"], ch, res["zp"])
        m = markers.get(r["Elemento"], ("o", "gray", 30))
        a1.scatter(r["c"], zc, marker=m[0], color=m[1], s=m[2], zorder=6,
                   label=None if r["Elemento"] in seen else r["Elemento"])
        seen.add(r["Elemento"])
    a1.set_ylabel("Elevación (msnm)")
    a1.grid(alpha=0.3)
    a1.legend(fontsize=7, ncol=3, loc="upper left")
    a1.set_title("Perfil de la línea de conducción por bombeo")
    a2.plot(ch, res["P"], color="red", label="Presión de operación (m.c.a.)")
    a2.plot(ch, res["Pst"], color="blue", ls="--", label="Presión estática (bomba parada)")
    a2.axhline(0, color="k", lw=0.8)
    if pn_m:
        a2.axhline(pn_m, color="m", ls=":", label="Clase de presión adoptada")
    a2.set_xlabel("Cadenamiento (m)")
    a2.set_ylabel("Presión (m.c.a.)")
    a2.grid(alpha=0.3)
    a2.legend(fontsize=7, ncol=3)
    fig.tight_layout()
    return fig


def fig_curva_masa(cum):
    fig, ax = plt.subplots(figsize=(7, 3))
    ax.bar(range(24), cum * 100, color="steelblue")
    ax.set_xlabel("Hora del día")
    ax.set_ylabel("Volumen acumulado (% del volumen diario)")
    ax.grid(alpha=0.3)
    ax.set_title("Curva masa de regulación")
    fig.tight_layout()
    return fig


# ----------------------------------------------------------------------------
# ESQUEMAS (no a escala; las cotas rotuladas SÍ son las calculadas)
# ----------------------------------------------------------------------------
from matplotlib.patches import Rectangle

C_AGUA, C_AGUA_U, C_TERR, C_SUELO = "#bcd9ec", "#4a90c4", "#6b4f2a", "#eadfca"


def _etiquetas(ax, items, x_line, x_text, yr, fs=6.5):
    """Rotula niveles evitando que los textos se encimen."""
    items = sorted(items, key=lambda t: -t[0])
    gap = 0.06 * yr
    ypos = []
    for y, _, _ in items:
        ypos.append(y if not ypos else min(y, ypos[-1] - gap))
    for (y, t, c), yp in zip(items, ypos):
        ax.plot([x_line, x_text - 0.2], [y, yp], color=c, lw=0.6)
        ax.text(x_text, yp, t, va="center", ha="left", fontsize=fs, color=c)


def draw_pozo(ax, d, fs=6.5, xmax=20):
    top, bot = d["z_t"], d["z_bomba"]
    yr = max(top - bot, 1.0)
    ax.set_xlim(0, xmax)
    ax.set_ylim(bot - 0.12 * yr, top + 0.12 * yr)
    ax.axis("off")
    ax.fill_between([0, 6.8], [top, top], [bot - 0.12 * yr] * 2, color=C_SUELO)
    ax.plot([0, 6.8], [top, top], color=C_TERR, lw=2)
    ax.add_patch(Rectangle((4, bot), 2, top - bot + 0.04 * yr, fc="white", ec="k", lw=1.2, zorder=2))
    ax.add_patch(Rectangle((4, bot), 2, d["z_ne"] - bot, fc=C_AGUA, ec="none", zorder=3))
    ax.add_patch(Rectangle((4, bot), 2, d["z_nd"] - bot, fc=C_AGUA_U, ec="none", zorder=3))
    ax.add_patch(Rectangle((4.35, d["z_succ"]), 1.3, 0.10 * yr, fc="gray", ec="k", zorder=4))
    ax.plot([5, 5], [d["z_succ"] + 0.10 * yr, top + 0.04 * yr], color="k", lw=2.5, zorder=4)
    _etiquetas(ax, [(top, f"Terreno {top:.2f}", C_TERR),
                    (d["z_ne"], f"NE {d['z_ne']:.2f}", "#1f77b4"),
                    (d["z_nd"], f"ND {d['z_nd']:.2f}", "#08519c"),
                    (d["z_nda"], f"ND + abatim. {d['z_nda']:.2f}", "#08306b"),
                    (d["z_succ"], f"Succión bomba {d['z_succ']:.2f}", "k"),
                    (bot, f"Colocación (c/holgura) {bot:.2f}", "gray")], 6, 7.2, yr, fs)
    ax.set_title("Pozo profundo (corte)", fontsize=fs + 2.5)


def draw_carcamo(ax, d, fs=6.5, xmax=20):
    top, bot = d["z_t"], d["z_fondo"]
    yr = max(top - bot, 1.0)
    ax.set_xlim(0, xmax)
    ax.set_ylim(bot - 0.12 * yr, top + 0.15 * yr)
    ax.axis("off")
    ax.fill_between([0, 9.9], [bot - 0.12 * yr] * 2, [bot] * 2, color=C_SUELO)
    ax.fill_between([0, 3], [top] * 2, [bot - 0.12 * yr] * 2, color=C_SUELO)
    ax.fill_between([9, 9.9], [top] * 2, [bot - 0.12 * yr] * 2, color=C_SUELO)
    ax.plot([0, 3], [top, top], color=C_TERR, lw=2)
    ax.plot([9, 9.9], [top, top], color=C_TERR, lw=2)
    ax.add_patch(Rectangle((3, bot), 6, top - bot + 0.05 * yr, fc="white", ec="k", lw=1.5, zorder=2))
    ax.add_patch(Rectangle((3, bot), 6, d["namino"] - bot, fc=C_AGUA, ec="none", zorder=3))
    ax.add_patch(Rectangle((3, d["namino"]), 6, d["namo"] - d["namino"], fc=C_AGUA_U, ec="none", zorder=3))
    ax.add_patch(Rectangle((5.2, bot + 0.01 * yr), 1.2, (d["z_succ"] - bot) + 0.10 * yr, fc="gray", ec="k", zorder=4))
    ax.plot([5.8, 5.8], [d["z_succ"] + 0.10 * yr, top + 0.05 * yr], color="k", lw=2.5, zorder=4)
    if fs < 8:
        ax.text(6.0, (d["namino"] + d["namo"]) / 2, "Vol. útil", fontsize=fs - 0.5, ha="left", va="center",
                color="white", zorder=5)
    _etiquetas(ax, [(top, f"Terreno {top:.2f}", C_TERR),
                    (d["name"], f"NAME {d['name']:.2f}", "#b30000"),
                    (d["namo"], f"NAMO (arranque) {d['namo']:.2f}", "#08519c"),
                    (d["namino"], f"NAMINO (paro) {d['namino']:.2f}", "#08306b"),
                    (d["z_succ"], f"Succión bomba {d['z_succ']:.2f}", "k"),
                    (bot, f"Fondo {bot:.2f}", "gray")], 9, 10.2, yr, fs)
    ax.set_title("Cárcamo de bombeo (corte)", fontsize=fs + 2.5)


def draw_tanque(ax, s, fs=6.5, xmax=20, lab_der=False):
    z_t, fo, no, co = s["z_t"], s["fondo"], s["namo"], s["corona"]
    lo = min(fo, z_t)
    yr = max(co - lo, 1.0)
    ax.set_xlim(0, xmax)
    ax.set_ylim(lo - 0.18 * yr, co + 0.22 * yr)
    ax.axis("off")
    ax.fill_between([0, 7.5], [z_t] * 2, [lo - 0.18 * yr] * 2, color=C_SUELO)
    ax.plot([0, 7.5], [z_t, z_t], color=C_TERR, lw=2)
    extra = []
    elevado = fo > z_t + 0.02 * yr or (s.get("base_h", 0) > 0 and fo > z_t + 1e-6)
    if elevado and s.get("base_tipo") == "Firme de concreto":
        ax.add_patch(Rectangle((2.6, z_t), 5.0, fo - z_t, fc="#bdbdbd", ec="k", lw=1.0, zorder=2))
        if lab_der:
            extra.append(((z_t + fo) / 2, f"Firme de concreto {fo - z_t:.2f} m", "#555555"))
        else:
            ax.text(0.2, (z_t + fo) / 2, f"Firme de concreto\n{fo - z_t:.2f} m", fontsize=fs, va="center")
    elif elevado:
        for x in (3.4, 6.6):
            ax.plot([x, x], [z_t, fo], color="k", lw=2.5)
        ax.plot([3.4, 6.6], [z_t, fo], color="k", lw=0.8)
        ax.plot([6.6, 3.4], [z_t, fo], color="k", lw=0.8)
        if lab_der:
            extra.append(((z_t + fo) / 2, f"{s.get('base_tipo') or 'Torre'} {fo - z_t:.2f} m", "#555555"))
        else:
            ax.text(0.2, (z_t + fo) / 2, f"{s.get('base_tipo') or 'Torre'}\n{fo - z_t:.2f} m", fontsize=fs, va="center")
    ax.add_patch(Rectangle((3, fo), 4, co - fo, fc="white", ec="k", lw=1.5, zorder=3))
    ax.add_patch(Rectangle((3, fo), 4, no - fo, fc=C_AGUA_U, ec="none", zorder=4))
    if s.get("recibe"):
        ax.annotate("Llegada\nde la línea", xy=(3.0, no + 0.05 * yr), xytext=(0.1, co + 0.10 * yr),
                    fontsize=fs, arrowprops=dict(arrowstyle="->", color="navy"), color="navy")
    items = [(co, f"Corona {co:.2f}", "#b30000"), (no, f"NAMO {no:.2f}", "#08519c"),
             (fo, f"Fondo = NAMIN {fo:.2f}", "#08306b")]
    if abs(z_t - fo) > 0.05 * yr:
        items.append((z_t, f"Terreno {z_t:.2f}", C_TERR))
    items += extra
    _etiquetas(ax, items, 7, 8.2, yr, fs)
    ax.set_title(s["nombre"] + (f"\n(×{s['n']} unidades, mismo nivel)" if s.get("n", 1) > 1 else ""), fontsize=fs + 2.5)


def fig_planta(items, forma):
    """items: [(nombre, n_unidades, area_por_unidad_m2)] -> planta a escala."""
    fig, ax = plt.subplots(figsize=(8, 2.8))
    x, ymax = 0.0, 0.0
    for nombre, n, A in items:
        if forma == "Circular":
            w = h = math.sqrt(4 * A / math.pi)
            lab = f"{n} × Ø {w:.2f} m"
        else:
            h = math.sqrt(A / 1.5)
            w = 1.5 * h
            lab = f"{n} × {w:.2f} × {h:.2f} m"
        x0 = x
        for _ in range(n):
            if forma == "Circular":
                ax.add_patch(plt.Circle((x + w / 2, h / 2), w / 2, fc=C_AGUA, ec="k"))
            else:
                ax.add_patch(Rectangle((x, 0), w, h, fc=C_AGUA, ec="k"))
            x += w * 1.08
        ax.text((x0 + x - 0.08 * w) / 2, -0.15 * h, f"{nombre}\n{lab}", ha="center", va="top", fontsize=7)
        ymax = max(ymax, h)
        x += 0.6 * w
    ax.set_xlim(-0.05 * x, x)
    ax.set_ylim(-0.8 * ymax, 1.15 * ymax)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("Planta de las estructuras (a escala)", fontsize=9)
    fig.tight_layout()
    return fig


def fig_pozo(d):
    fig, ax = plt.subplots(figsize=(4.4, 4.0))
    draw_pozo(ax, d)
    fig.tight_layout()
    return fig


def fig_carcamo(d):
    fig, ax = plt.subplots(figsize=(4.4, 4.0))
    draw_carcamo(ax, d)
    fig.tight_layout()
    return fig


def fig_tanques(estr):
    n = len(estr)
    fig, axs = plt.subplots(1, n, figsize=(4.4 * n, 4.0))
    axs = np.atleast_1d(axs)
    for ax, s_ in zip(axs, estr):
        draw_tanque(ax, s_)
    fig.tight_layout()
    return fig


def fig_descarga(destino, llega_en, h, cist_piso=False):
    """Croquis pequeño: qué altura capturar en el punto de descarga."""
    fig, ax = plt.subplots(figsize=(4.2, 2.8))
    ax.set_xlim(0, 10)
    ax.set_ylim(-2.4, 7.6)
    ax.axis("off")
    ax.fill_between([0, 10], [0, 0], [-2.4, -2.4], color=C_SUELO)
    ax.plot([0, 10], [0, 0], color=C_TERR, lw=2)
    ax.text(6.9, -0.75, "Terreno final", fontsize=6.5, color=C_TERR)
    elevado = llega_en == "Tanque elevado"
    cist = llega_en == "Cisterna"
    if elevado:
        for x in (3.8, 6.2):
            ax.plot([x, x], [0, 3.4], color="k", lw=2.5)
        ax.add_patch(Rectangle((3.2, 3.4), 3.6, 2.4, fc="white", ec="k", lw=1.5))
        ax.add_patch(Rectangle((3.2, 3.4), 3.6, 2.0, fc=C_AGUA_U, ec="none"))
        y_n, lab = 5.4, "h = torre + tirante"
        ax.annotate("", xy=(2.6, 3.4), xytext=(2.6, 0), arrowprops=dict(arrowstyle="<->", color="gray"))
        ax.text(0.2, 1.6, "torre", fontsize=6.5, color="gray")
        y_in = 5.9
    elif cist and not cist_piso:
        ax.add_patch(Rectangle((3.2, -1.9), 3.6, 3.0, fc="white", ec="k", lw=1.5, zorder=3))
        ax.add_patch(Rectangle((3.2, -1.9), 3.6, 2.5, fc=C_AGUA_U, ec="none", zorder=4))
        y_n, lab = 0.6, "h (≈ 0 si el agua\nllega al nivel del terreno)"
        y_in = 1.3
    else:
        ax.add_patch(Rectangle((3.2, 0), 3.6, 3.4, fc="white", ec="k", lw=1.5, zorder=3))
        ax.add_patch(Rectangle((3.2, 0), 3.6, 2.9, fc=C_AGUA_U, ec="none", zorder=4))
        y_n, lab = 2.9, "h ≈ tirante máx."
        y_in = 3.7
    ax.plot([0.2, 2.2, 2.2, 3.6], [-0.6, -0.6, y_in, y_in], color="navy", lw=2, zorder=5)
    ax.annotate("", xy=(3.6, y_in), xytext=(3.3, y_in), arrowprops=dict(arrowstyle="->", color="navy"), zorder=5)
    ax.text(0.2, -1.3, "Línea de conducción", fontsize=6.5, color="navy")
    ax.plot([6.8, 8.4], [y_n, y_n], color="red", ls="--", lw=1)
    ax.annotate("", xy=(8.2, y_n), xytext=(8.2, 0), arrowprops=dict(arrowstyle="<->", color="red", lw=1.4))
    ax.text(8.35, y_n / 2 if y_n > 1 else y_n + 0.5, lab, fontsize=7, color="red", va="center")
    ax.text(6.9, y_n + 0.15, "NAMO", fontsize=6.5, color="red")
    ax.set_title(f"Altura de descarga h  (valor actual: {h:.1f} m)", fontsize=8)
    fig.tight_layout()
    return fig


def fig_sistema(fd, res, estr):
    n = 2 + len(estr)
    fig, axs = plt.subplots(1, n, figsize=(4.0 * n + 2.0, 4.8),
                            gridspec_kw={"width_ratios": [1.0, 1.9] + [1.0] * len(estr)})
    (draw_pozo if fd["tipo"] == "pozo" else draw_carcamo)(axs[0], fd)
    ax = axs[1]
    ch = res["ch"]
    ax.plot(ch, res["zt"], color=C_TERR, ls="--", lw=1, label="Terreno")
    ax.plot(ch, res["zp"], color="navy", lw=1.8, label="Tubería")
    ax.plot(ch, res["LE"], color="red", ls="--", lw=1.3, label="Línea de energía")
    i_hi = int(np.argmax(res["zp"]))
    ax.scatter([ch[0], ch[i_hi], ch[-1]], [res["zp"][0], res["zp"][i_hi], res["zp"][-1]], color="navy", zorder=5)
    ax.annotate(f"Inicio tubería {res['zp'][0]:.2f}\nLE en bomba {res['LE'][0]:.2f}", (ch[0], res["LE"][0]),
                textcoords="offset points", xytext=(6, -38), fontsize=7, color="red")
    ax.annotate(f"Punto más alto {res['zp'][i_hi]:.2f}\n(cad. {ch[i_hi]:.0f} m)", (ch[i_hi], res["zp"][i_hi]),
                textcoords="offset points", xytext=(-12, 14), fontsize=7, color="navy", ha="right")
    ax.annotate(f"Fin tubería {res['zp'][-1]:.2f}\nLE llegada {res['LE'][-1]:.2f}", (ch[-1], res["zp"][-1]),
                textcoords="offset points", xytext=(-100, -46), fontsize=7, color="navy")
    ax.set_xlabel("Cadenamiento (m)")
    ax.set_ylabel("Cota (msnm)")
    ax.grid(alpha=0.3)
    ax.legend(fontsize=7, loc="lower right")
    ax.set_title("Línea de conducción", fontsize=9)
    for a_, s_ in zip(axs[2:], estr):
        draw_tanque(a_, s_)
    fig.tight_layout()
    return fig


EXE_NAVY, EXE_TEAL, EXE_ORANGE, EXE_RED, EXE_GRAY = "#0B3C5D", "#1D7A8C", "#E07A1F", "#C0392B", "#5F6B76"


def fig_sistema_exec(fd, res, estr, info, size=(13.333, 7.5), fscale=1.0, kpi=True, titulo=True,
                     show_alto=False, show_llegada=False):
    """Versión ejecutiva del esquema general (para diapositivas): texto grande, paleta sobria,
    encabezados por bloque y franja de datos clave."""
    fs = 10.0 * fscale
    W = size[0]
    fig = plt.figure(figsize=size)
    left, right, gap = 0.035, 0.985, 0.016
    top_y = 0.80 if titulo else 0.90
    bot_y = 0.215 if kpi else 0.08
    ratios = ([1.5, 2.0] + [1.0] * len(estr)) if W >= 12 else ([1.4, 1.7] + [1.2] * len(estr))
    total = (right - left) - gap * (len(ratios) - 1)
    boxes, x = [], left
    for r in ratios:
        w = total * r / sum(ratios)
        boxes.append((x, w))
        x += w + gap
    axs = [fig.add_axes([x0, bot_y, w, top_y - bot_y]) for x0, w in boxes]
    heads = [f"1 · Fuente: {info['fuente']}", "2 · Línea de conducción"] + \
            [f"{3 + i} · {e['nombre']}" + (f" (×{e['n']})" if e.get("n", 1) > 1 else "") for i, e in enumerate(estr)]
    for (x0, w), h_ in zip(boxes, heads):
        fig.add_artist(Rectangle((x0, top_y + 0.008), w, 0.052, transform=fig.transFigure, fc=EXE_NAVY, ec="none"))
        h_ = h_.replace("Tanque de regulación (piso)", "Tanque de piso")
        cpl = max(int(w * W * 72 * 0.92 / (0.60 * 1.2 * fs)), 8)
        lines_ = textwrap.wrap(h_, cpl) or [h_]
        fsh_ = (1.2 if len(lines_) == 1 else 0.95) * fs
        if len(lines_) > 2:
            fsh_ = 0.8 * fs
            lines_ = textwrap.wrap(h_, max(int(w * W * 72 * 0.92 / (0.60 * fsh_)), 8)) or [h_]
        fig.text(x0 + w / 2, top_y + 0.034, "\n".join(lines_), ha="center", va="center", color="white",
                 fontsize=fsh_, fontweight="bold", linespacing=1.1)
    if titulo:
        fst = min(2.1 * fs, (right - left) * W * 72 * 0.97 / (0.60 * max(len(info["titulo"]), 1)))
        fig.text(left, 0.968, info["titulo"], fontsize=fst, fontweight="bold", color=EXE_NAVY, va="top")
        fig.text(left, 0.912, info["subtitulo"], fontsize=1.1 * fs, color=EXE_GRAY, va="top")

    # 1) fuente
    (draw_pozo if fd["tipo"] == "pozo" else draw_carcamo)(axs[0], fd, fs=fs, xmax=30)
    axs[0].set_title("")

    # 2) línea de conducción
    ax = axs[1]
    ch, zt, zp, LE = res["ch"], res["zt"], res["zp"], res["LE"]
    span = max(LE.max() - zp.min(), 1.0)
    lo, hi = zp.min() - 0.10 * span, LE.max() + 0.14 * span
    ax.set_ylim(lo, hi)
    ax.set_xlim(ch[0] - 0.02 * (ch[-1] - ch[0]), ch[-1] + 0.02 * (ch[-1] - ch[0]))
    ax.fill_between(ch, zt, lo, color=C_SUELO, zorder=1)
    ax.plot(ch, zt, color=C_TERR, lw=1.6, label="Terreno", zorder=2)
    ax.plot(ch, zp, color=EXE_NAVY, lw=3.2, label="Tubería", zorder=4)
    ax.plot(ch, LE, color=EXE_RED, lw=2.0, ls="--", label="Línea de energía", zorder=3)
    acc_ = res["accr"]
    for nm_, mk, col, sz, lab in [("Válvula de expulsión de aire", "*", EXE_ORANGE, 260, "Válvula de aire"),
                                  ("Válvula de desagüe (purga)", "s", EXE_TEAL, 90, "Desagüe")]:
        pts = acc_.loc[acc_["Elemento"] == nm_, "c"].to_numpy(float) if len(acc_) else []
        if len(pts):
            ax.scatter(pts, np.interp(pts, ch, zp), marker=mk, s=sz, color=col, edgecolor="k", lw=0.8,
                       zorder=6, label=lab)
    i_hi = int(np.argmax(zp))
    bb = dict(boxstyle="round,pad=0.25", fc="white", ec="#C9D3DC", alpha=0.95)
    ax.annotate(f"Bomba\nLE {LE[0]:.2f}", (ch[0], LE[0]), xytext=(10, 6), textcoords="offset points",
                fontsize=fs, color=EXE_RED, fontweight="bold", va="bottom", bbox=bb)
    if show_alto:
        ax.annotate(f"Punto más alto\n{zp[i_hi]:.2f} msnm (cad. {ch[i_hi]:.0f})", (ch[i_hi], zp[i_hi]),
                    xytext=(-70, -48), textcoords="offset points", fontsize=fs, color=EXE_NAVY, ha="right",
                    bbox=bb, arrowprops=dict(arrowstyle="-", color=EXE_NAVY, lw=0.8), zorder=7)
    if show_llegada:
        ax.annotate(f"Llegada\ntubería {zp[-1]:.2f}\nLE {LE[-1]:.2f}", (ch[-1], zp[-1]),
                    xytext=(-14, -66), textcoords="offset points", fontsize=fs, color=EXE_NAVY, ha="right",
                    bbox=bb, arrowprops=dict(arrowstyle="-", color=EXE_NAVY, lw=0.8), zorder=7)
    ax.text(0.99, 0.97, info["callout"], transform=ax.transAxes, ha="right", va="top", fontsize=fs,
            color=EXE_NAVY, bbox=dict(boxstyle="round,pad=0.45", fc="#F4F7FA", ec="#C9D3DC"), zorder=8)
    ax.set_xlabel("Cadenamiento (m)", fontsize=fs)
    ax.set_ylabel("Cota (msnm)", fontsize=fs)
    ax.tick_params(labelsize=0.9 * fs)
    ax.grid(alpha=0.25)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.legend(fontsize=0.9 * fs, loc="lower right", frameon=True, framealpha=0.95)

    # 3..n) almacenamiento
    for ax_, e in zip(axs[2:], estr):
        draw_tanque(ax_, e, fs=fs * (1.0 if len(estr) <= 2 else 0.85),
                    xmax=(30 if W >= 12 else 34) + (6 if len(estr) >= 3 else 0), lab_der=True)
        ax_.set_title("")

    # franja de datos clave
    if kpi and info.get("kpis"):
        axk = fig.add_axes([left, 0.035, right - left, 0.135])
        axk.set_xlim(0, 1)
        axk.set_ylim(0, 1)
        axk.axis("off")
        n = len(info["kpis"])
        wk = 1.0 / n
        box_in = (wk - 0.008) * (right - left) * W
        fsv = min([1.7 * fs] + [box_in * 72 * 0.86 / (0.58 * max(len(v_), 1)) for _, v_ in info["kpis"]])
        fsl = min([0.9 * fs] + [box_in * 72 * 0.88 / (0.56 * max(len(l_), 1)) for l_, _ in info["kpis"]])
        for i, (lab, val) in enumerate(info["kpis"]):
            x0 = i * wk + 0.004
            axk.add_patch(Rectangle((x0, 0.04), wk - 0.008, 0.92, fc="#F4F7FA", ec="#D5DEE6", lw=1))
            axk.add_patch(Rectangle((x0, 0.04), 0.005, 0.92, fc=EXE_TEAL, ec="none"))
            axk.text(x0 + 0.014, 0.74, lab, fontsize=fsl, color=EXE_GRAY, va="center", ha="left")
            axk.text(x0 + 0.014, 0.36, val, fontsize=fsv, color=EXE_NAVY, fontweight="bold", va="center",
                     ha="left")
    fig.text(right, 0.008, info.get("pie", ""), ha="right", va="bottom", fontsize=0.8 * fs, color=EXE_GRAY)
    return fig


# ----------------------------------------------------------------------------
# EXPORTACIÓN
# ----------------------------------------------------------------------------
def build_excel(sheets):
    from openpyxl.utils import get_column_letter
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for i, col in enumerate(df.columns, 1):
                w = max([len(str(col))] + [len(str(x)) for x in df[col].head(60)]) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(w, 55)
    return buf.getvalue()


def build_docx(resumen, alertas, figs, titulo, tablas=None, parrafos=None):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading(titulo, 0)
    doc.add_paragraph(f"Fecha de elaboración: {dt.date.today():%d/%m/%Y}. "
                      "Criterios de referencia: manuales MAPAS - CONAGUA. "
                      "Documento generado automáticamente; debe ser revisado por el proyectista responsable.")
    secciones = list(dict.fromkeys(resumen["Sección"]))
    for s in secciones:
        doc.add_heading(s, 1)
        sub = resumen[resumen["Sección"] == s]
        t = doc.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(["Concepto", "Valor", "Unidad"]):
            t.rows[0].cells[i].text = h
        for _, r in sub.iterrows():
            c = t.add_row().cells
            c[0].text, c[1].text, c[2].text = str(r["Concepto"]), str(r["Valor"]), str(r["Unidad"])
    if parrafos:
        doc.add_heading("Descripción de las estructuras de almacenamiento", 1)
        for t_ in parrafos:
            doc.add_paragraph(t_)
    for ttl, dfx in (tablas or []):
        doc.add_heading(ttl, 1)
        tb = doc.add_table(rows=1, cols=len(dfx.columns))
        tb.style = "Light Grid Accent 1"
        for i, h in enumerate(dfx.columns):
            tb.rows[0].cells[i].text = str(h)
        for _, r in dfx.iterrows():
            cc = tb.add_row().cells
            for i, v in enumerate(r):
                cc[i].text = str(v)
    doc.add_heading("Alertas y observaciones", 1)
    for a in alertas or ["Sin alertas."]:
        doc.add_paragraph(a, style="List Bullet")
    doc.add_heading("Figuras", 1)
    for cap, fig in figs:
        b = io.BytesIO()
        fig.savefig(b, format="png", dpi=150)
        b.seek(0)
        doc.add_picture(b, width=Inches(6.3))
        doc.add_paragraph(cap)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ============================================================================
# INTERFAZ
# ============================================================================
st.set_page_config(page_title="Memoria de línea de bombeo - MAPAS", layout="wide")
st.title("Memoria técnica descriptiva · Línea de bombeo (criterios MAPAS-CONAGUA)")
st.caption("Valores precargados = referencia editable. Verifica contra el manual vigente.")

RES = []          # (sección, concepto, valor, unidad)
ALERTAS = []
SHEETS = {}
FIGS = []
TABLAS_DOC = []     # (título, DataFrame) extra para la memoria Word
PARRAFOS_DOC = []   # párrafos descriptivos para la memoria Word


def add(sec, name, val, unit=""):
    RES.append((sec, name, round(val, 3) if isinstance(val, (float, np.floating)) else val, unit))


tabs = st.tabs(["1 Datos", "2 Perfil", "3 Material y diámetro", "4 Accesorios",
                "5 Hidráulica y bomba", "6 Fuente (pozo/cárcamo)", "7 Tanques", "8 Exportar"])

# ------------------------------------------------------------------ 1 DATOS
with tabs[0]:
    c1, c2, c3 = st.columns(3)
    with c1:
        proyecto = st.text_input("Nombre del proyecto", "Línea de bombeo - Ejemplo")
        fuente = st.selectbox("Tipo de fuente de abastecimiento", list(EQUIPOS))
        equipo = st.selectbox("Tipo de equipo de bombeo", EQUIPOS[fuente])
    with c2:
        T = st.number_input("Horas de bombeo al día", 1.0, 24.0, 20.0, 0.5)
        h0 = st.number_input("Hora de arranque (0-23)", 0, 23, 4)
        eta_b = st.number_input("Eficiencia de la bomba", 0.30, 0.90, ETA_B[equipo], 0.01)
        eta_m = st.number_input("Eficiencia del motor", 0.60, 0.97, 0.90, 0.01)
        fs = st.number_input("Factor de servicio del motor", 1.0, 1.5, 1.15, 0.05)
    with c3:
        destino = st.selectbox("Destino", ["Tanque de regulación superficial (piso)", "Tanque elevado",
                                           "Cisterna", "Cisterna + tanque elevado"] + DEST_NUEVOS)
        llega_en = {"Tanque elevado": "Tanque elevado", "Cisterna": "Cisterna"}.get(destino, "Tanque de piso")
        if destino == "Cisterna + tanque elevado":
            llega_en = st.selectbox("La línea de conducción descarga en:", ["Tanque elevado", "Cisterna"])
        elif destino in DEST_NUEVOS:
            llega_en = st.selectbox("La línea de conducción descarga en:", DEST_ESTR[destino], key="llega_nuevo")
        cist_piso = False
        if "Cisterna" in destino:
            cist_tipo = st.selectbox("Tipo de construcción de la cisterna",
                                     ["Enterrada / semienterrada (con excavación)",
                                      "Sobre el terreno, tipo tanque de piso (sin excavación)"])
            cist_piso = cist_tipo.startswith("Sobre")
            add("Datos", "Tipo de construcción de la cisterna", cist_tipo)
        cist_recibe_piso = cist_piso and llega_en == "Cisterna"
        st.session_state.setdefault("h_llegada_in", 4.0)
        d_tir, caida = 3.5, 0.0
        desg = st.checkbox("Desglosar la altura de descarga h (base/torre + tirante + caída)", value=False,
                           key="h_desg_on")
        if desg:
            _lbl = {"Tanque elevado": "Altura de la torre: fondo del tanque sobre el terreno (m)",
                    "Tanque de piso": "Altura de la base (firme o estructura) sobre el terreno (m)",
                    "Cisterna": "Fondo de la cisterna respecto al terreno (m; negativo = enterrada)"}[llega_en]
            _def = {"Tanque elevado": 12.0, "Tanque de piso": 0.0, "Cisterna": 0.0 if cist_piso else -3.5}[llega_en]
            d_base = st.number_input(_lbl, -20.0, 60.0, _def, 0.05, key=f"h_d_base_{llega_en}_{int(cist_piso)}")
            d_tir = st.number_input("Tirante máximo del agua (m)", 0.5, 20.0, 3.5, 0.05, key="h_d_tir")
            _tll = st.selectbox("Tipo de llegada", ["Sumergida / al nivel del agua",
                                                    "Descarga libre por arriba (caída sobre el NAMO)"], key="h_d_tll")
            if _tll.startswith("Descarga"):
                caida = st.number_input("Caída libre de la boca de descarga sobre el NAMO (m)", 0.0, 5.0, 0.30, 0.05,
                                        key="h_d_caida")
            st.session_state["h_llegada_in"] = round(d_base + d_tir + caida, 3)
            add("Datos", "Desglose de h (base/torre + tirante + caída)",
                f"{d_base:.2f} + {d_tir:.2f} + {caida:.2f} = {d_base + d_tir + caida:.2f} m")
        h_llegada = st.number_input(
            "Altura de descarga h: nivel MÁXIMO del agua sobre el terreno final (m) — ver croquis",
            min_value=-20.0, max_value=80.0, step=0.05, key="h_llegada_in", format="%.2f", disabled=desg)
        cover = st.number_input("Profundidad de colocación de tubería (m)", 0.0, 5.0, 1.0, 0.1)
        v_min = st.number_input("Velocidad mínima recomendada (m/s)", 0.1, 2.0, 0.6, 0.1)
        v_max = st.number_input("Velocidad máxima recomendada (m/s)", 0.5, 5.0, 2.0, 0.1)

    st.markdown("---")
    cA, cB = st.columns([1.7, 1])
    with cA:
        st.markdown("#### Caudales de diseño")
        MODOS = ["Población y dotación", "Conozco Qmed", "Conozco el caudal de bombeo Qb"]
        modo_q = st.radio("¿Cómo definir el caudal?", MODOS, horizontal=True)
        modo = MODOS.index(modo_q)
        k1, k2 = st.columns(2)
        dot = pob = None
        if modo == 0:
            dot = k1.number_input("Dotación (l/hab/día)", 20, 1000, 150, 5)
            pob = k2.number_input("Población de proyecto (hab)", 1, 5000000, 6000, 100)
            Qmed_in = pob * dot / 86400.0
        elif modo == 1:
            Qmed_in = k1.number_input("Qmed (l/s)", 0.1, 5000.0, 10.4, 0.1)
        else:
            Qb_in = k1.number_input("Qb, caudal a bombear (l/s)", 0.1, 5000.0, 20.0, 0.5)
        k3, k4 = st.columns(2)
        CVd = k3.number_input("Coeficiente de variación diaria CVd", 1.0, 2.5, 1.20, 0.01)
        CVh = k4.number_input("Coeficiente de variación horaria CVh", 1.0, 4.0, 1.50, 0.01)
        if modo in (0, 1):
            Qmed = Qmed_in
            Qmd = Qmed * CVd
            Qlps = Qmd * 24.0 / T
        else:
            Qlps = Qb_in
            Qmd = Qlps * T / 24.0
            Qmed = Qmd / CVd
        Qmh = Qmd * CVh          # MAPAS: Qmh = Qmd × CVh
        Qlps = max(Qlps, 0.1)
        tq_q = pd.DataFrame([
            ["Qmed", round(Qmed, 3), "Gasto medio diario (Pob × Dot / 86 400)" if modo == 0 else "Gasto medio diario"],
            ["Qmd = Qmed × CVd", round(Qmd, 3),
             "Obra de captación, línea de conducción (hasta el tanque) y tanque de regulación"],
            ["Qmh = Qmd × CVh", round(Qmh, 3), "Línea de alimentación a la red y red de distribución"],
            ["Qb = Qmd × 24 / T", round(Qlps, 3), "Caudal de bombeo: equipo, columna/succión y línea de conducción"],
        ], columns=["Caudal", "l/s", "Se usa para diseñar"])
        st.dataframe(tq_q, width="stretch", hide_index=True)
        if Qlps > 0 and T < 24:
            st.caption(f"Bombeando {T:g} h/día el equipo debe entregar Qb = Qmd·24/T = {Qlps:.2f} l/s para cubrir Qmd.")
    with cB:
        st.markdown("#### ¿Qué altura capturar en la descarga?")
        figd = fig_descarga(destino, llega_en, h_llegada, cist_piso)
        st.pyplot(figd)
        plt.close(figd)
        st.caption("h = distancia vertical entre el terreno en el último punto del perfil y el nivel máximo del "
                   "agua (NAMO) al que descarga la línea. Piso: ≈ tirante; elevado: torre + tirante; "
                   "cisterna enterrada: ≈ 0 si el agua llega al nivel del terreno; cisterna sobre el terreno (tipo piso, "
                   "sin excavación): ≈ tirante.")
    if modo == 0:
        add("Datos", "Dotación", dot, "l/hab/día")
        add("Datos", "Población de proyecto", pob, "hab")
    add("Datos", "Qmed", float(Qmed), "l/s")
    add("Datos", "CVd / CVh", f"{CVd} / {CVh}")
    add("Datos", "Qmd (captación, conducción, regulación)", float(Qmd), "l/s")
    add("Datos", "Qmh (alimentación a red y red)", float(Qmh), "l/s")
Q = Qlps / 1000.0

# ------------------------------------------------------------------ 2 PERFIL
with tabs[1]:
    st.write("Sube CSV/Excel con columnas **Cadenamiento** (m, o formato 0+500) y **Cota** (msnm), "
             "desde la fuente hasta el tanque.")
    plantilla = demo_profile().rename(columns={"ch": "Cadenamiento", "z": "Cota"})
    st.download_button("Descargar plantilla CSV", plantilla.to_csv(index=False).encode("utf-8"),
                       "plantilla_perfil.csv")
    up = st.file_uploader("Perfil del TERRENO", type=["csv", "xlsx", "xls"], key="up_terreno")
    terr = read_profile(up, "terreno", "prf_terr") if up is not None else None
    if terr is None:
        st.info("Usando perfil de ejemplo hasta que subas un archivo válido.")
        terr = demo_profile()
    with st.expander("Opcional: perfil de la TUBERÍA (si ya lo tienes en el mismo formato)"):
        st.caption("Mismo formato que el terreno: Cadenamiento y Cota de tubería (rasante/eje). "
                   "Los cadenamientos pueden ser distintos a los del terreno; se interpolan. "
                   "Si no lo cargas, la tubería = terreno − profundidad de colocación.")
        up_p = st.file_uploader("Perfil de la TUBERÍA", type=["csv", "xlsx", "xls"], key="up_tuberia")
        pp = read_profile(up_p, "tuberia", "prf_tub") if up_p is not None else None
    if pp is not None:
        lo_, hi_ = float(terr["ch"].min()), float(terr["ch"].max())
        ch_all = np.union1d(terr["ch"].to_numpy(float), pp["ch"].to_numpy(float))
        ch_all = ch_all[(ch_all >= lo_) & (ch_all <= hi_)]
        prof = pd.DataFrame({"ch": ch_all,
                             "z": np.round(np.interp(ch_all, terr["ch"], terr["z"]), 3),
                             "zp": np.round(np.interp(ch_all, pp["ch"], pp["z"]), 3)})
        if pp["ch"].min() > lo_ + 0.5 or pp["ch"].max() < hi_ - 0.5:
            st.warning("El perfil de tubería no cubre todo el cadenamiento del terreno; "
                       "los extremos se completan con el valor más cercano.")
        cob = prof["z"] - prof["zp"]
        st.write(f"Tubería cargada por el usuario. Profundidad (terreno − tubería): "
                 f"mín **{cob.min():.2f} m**, máx **{cob.max():.2f} m**.")
        if cob.min() < 0:
            st.error("En algún punto la tubería queda por encima del terreno; revisa el archivo.")
            ALERTAS.append("El perfil de tubería cargado queda por encima del terreno en algún tramo.")
        add("Datos", "Perfil de tubería", "Cargado por el usuario")
    else:
        prof = terr.copy()
        prof["zp"] = prof["z"] - cover
        add("Datos", "Perfil de tubería", f"Terreno − {cover} m")
    st.dataframe(prof.rename(columns={"ch": "Cadenamiento (m)", "z": "Cota terreno (msnm)",
                                      "zp": "Cota tubería (msnm)"}), height=200)
    fp, axp = plt.subplots(figsize=(11, 3))
    axp.plot(prof["ch"], prof["z"], color="saddlebrown", label="Terreno")
    axp.plot(prof["ch"], prof["zp"], color="navy", ls="--", label="Tubería")
    axp.fill_between(prof["ch"], prof["z"], prof["z"].min() - 5, color="burlywood", alpha=0.3)
    axp.set_xlabel("Cadenamiento (m)")
    axp.set_ylabel("Cota (msnm)")
    axp.grid(alpha=0.3)
    axp.legend(fontsize=8)
    st.pyplot(fp)
    plt.close(fp)

z_ground0 = float(prof["z"].iloc[0])
z_end = float(prof["z"].iloc[-1])
z_dest = z_end + h_llegada

# ------------------------------------------------------------------ 3 MATERIAL / DIÁMETRO
with tabs[2]:
    c1, c2 = st.columns(2)
    with c1:
        mat_name = st.selectbox("Material", list(MATERIALES) + ["Otro (manual)"])
        base = MATERIALES.get(mat_name, dict(eps=0.1, C=120, f_di=1.0, a=900))
        if mat_name == "Otro (manual)":
            mat_name = st.text_input("Nombre del material", "Material especial")
            eps = st.number_input("Rugosidad absoluta ε (mm)", 0.0001, 5.0, 0.1, format="%.4f")
            Chw = st.number_input("C de Hazen-Williams", 60, 160, 120)
            f_di = st.number_input("Relación DI/DN", 0.5, 1.2, 1.0, 0.01)
            a_cel = st.number_input("Celeridad onda (m/s)", 100, 1500, 900)
        else:
            eps = st.number_input("Rugosidad absoluta ε (mm) [editable]", 0.0001, 5.0, float(base["eps"]),
                                  format="%.4f")
            Chw = st.number_input("C de Hazen-Williams [editable]", 60, 160, int(base["C"]))
            f_di = st.number_input("Relación DI/DN aprox. [usa DI de catálogo si lo tienes]", 0.5, 1.2,
                                   float(base["f_di"]), 0.01)
            a_cel = base["a"]
        metodo = st.radio("Ecuación de fricción", ["Darcy-Weisbach (Swamee-Jain)", "Hazen-Williams"])
    # tabla comparativa: resultados con accesorios actuales se calculan después;
    # aquí se usa la tabla de accesorios guardada en sesión.
    if "acc_v" not in st.session_state:
        st.session_state.acc_v = 0
    acc_state = st.session_state.get("acc", pd.DataFrame(
        columns=["Cadenamiento (m)", "Elemento", "Cantidad", "K", "Diámetro (mm)"]))
    z_low_prelim = z_ground0 - 10   # solo para comparativo; el real se calcula en pestaña 6
    cand = []
    Dteo = math.sqrt(4 * Q / (math.pi * 1.2)) * 1000
    for dn in DN_COMERCIALES:
        di = dn * f_di / 1000
        if di > 1.0 and dn > 600 and Q < 0.05:
            continue
        r = hidraulica(prof, cover, di, Q, eps, Chw, metodo, acc_state, z_dest)
        ok = v_min <= r["v"] <= v_max
        cand.append(dict(DN_mm=dn, DI_mm=round(di * 1000, 1), V_ms=round(r["v"], 2),
                         hf_m=round(r["hf"], 2), hl_m=round(r["hl"], 2),
                         J_m_km=round(r["hf"] / max(r["L"], 1) * 1000, 2),
                         Cumple_v="SÍ" if ok else "NO"))
    cand = pd.DataFrame(cand)
    with c2:
        st.write(f"Diámetro teórico (v≈1.2 m/s): **{Dteo:.0f} mm**")
        st.dataframe(cand, height=300)
        ok_list = cand[cand["Cumple_v"] == "SÍ"]["DN_mm"].tolist()
        dn_default = ok_list[len(ok_list) // 2] if ok_list else 200
        DN = st.selectbox("Diámetro nominal propuesto (mm)", DN_COMERCIALES,
                          index=DN_COMERCIALES.index(dn_default))
        di_ov = st.number_input("Diámetro interior real (mm) - 0 = DN × relación", 0.0, 2000.0, 0.0)
D = (di_ov if di_ov > 0 else DN * f_di) / 1000.0
v_line = Q / (math.pi * D ** 2 / 4)
with tabs[2]:
    if v_line < v_min:
        st.error(f"Velocidad {v_line:.2f} m/s < mínima {v_min}. Reduce el diámetro.")
        ALERTAS.append(f"Velocidad en línea {v_line:.2f} m/s menor a la mínima ({v_min} m/s).")
    elif v_line > v_max:
        st.error(f"Velocidad {v_line:.2f} m/s > máxima {v_max}. Aumenta el diámetro.")
        ALERTAS.append(f"Velocidad en línea {v_line:.2f} m/s mayor a la máxima ({v_max} m/s).")
    else:
        st.success(f"Velocidad {v_line:.2f} m/s: cumple ({v_min}-{v_max} m/s).")
    SHEETS["Diametros_candidatos"] = cand

# ------------------------------------------------------------------ 6 FUENTE (se calcula antes de 5 por dependencias)
# Diámetros de succión / descarga de bomba
D_succ = commercial_up(math.sqrt(4 * Q / (math.pi * 1.5)) * 1000)      # v ≤ 1.5 m/s
D_desc = commercial_up(math.sqrt(4 * Q / (math.pi * 3.0)) * 1000)      # v ≤ 3.0 m/s

with tabs[5]:
    st.subheader(fuente)
    if fuente == "Pozo profundo":
        c1, c2 = st.columns(2)
        with c1:
            ne = st.number_input("Profundidad del nivel estático NE (m bajo terreno)", 0.0, 800.0, 40.0, 1.0)
            nd = st.number_input("Profundidad del nivel dinámico ND a Qb (m bajo terreno)", 0.0, 800.0, 55.0, 1.0)
            abat = st.number_input("Abatimiento adicional esperado (m)", 0.0, 100.0, 5.0, 1.0)
        with c2:
            s_min = st.number_input("Sumergencia mínima sobre succión (m)", 0.5, 20.0, 3.0, 0.5)
            holg = st.number_input("Holgura sobre fondo de pozo / tazones (m)", 0.0, 20.0, 3.0, 0.5)
            d_col = st.number_input("Diámetro interior columna (mm)", 50.0, 600.0,
                                    float(commercial_up(math.sqrt(4 * Q / (math.pi * 2.0)) * 1000)))
        prof_bomba = nd + abat + s_min + holg
        h_col = hf_pipe(Q, d_col / 1000, prof_bomba, 0.15, 120, metodo) + 0.5   # +0.5 m cabezal/codo
        z_low = z_ground0 - nd
        z_high = z_ground0 - ne
        v_col = Q / (math.pi * (d_col / 1000) ** 2 / 4)
        st.write(f"Profundidad de colocación de bomba: **{prof_bomba:.1f} m**; pérdida en columna: "
                 f"**{h_col:.2f} m**; v columna = {v_col:.2f} m/s")
        if v_col > 2.5:
            ALERTAS.append(f"Velocidad en columna {v_col:.2f} m/s alta; aumentar diámetro de columna.")
        add("Pozo profundo", "Nivel estático (profundidad)", ne, "m")
        add("Pozo profundo", "Nivel dinámico (profundidad)", nd, "m")
        add("Pozo profundo", "Profundidad de colocación de bomba", prof_bomba, "m")
        add("Pozo profundo", "Pérdida en columna", h_col, "m")
        add("Pozo profundo", "Diámetro de columna", d_col, "mm")
        z_t = z_ground0
        FUENTE_D = dict(tipo="pozo", z_t=z_t, z_ne=z_t - ne, z_nd=z_t - nd, z_nda=z_t - nd - abat,
                        z_succ=z_t - (nd + abat + s_min), z_bomba=z_t - prof_bomba)
        NIV_FUENTE = [
            ("Terreno en el pozo", z_t, "Cota de terreno (inicio del perfil)"),
            ("Nivel estático (NE)", FUENTE_D["z_ne"], f"{ne:.1f} m bajo terreno"),
            ("Nivel dinámico (ND) a Qb", FUENTE_D["z_nd"], f"{nd:.1f} m bajo terreno"),
            ("ND + abatimiento adicional", FUENTE_D["z_nda"], f"{nd + abat:.1f} m bajo terreno"),
            ("Succión de la bomba", FUENTE_D["z_succ"], f"ND + abatim. + sumergencia = {nd + abat + s_min:.1f} m"),
            ("Profundidad de colocación (incl. holgura)", FUENTE_D["z_bomba"], f"{prof_bomba:.1f} m bajo terreno"),
        ]
        cf1, cf2 = st.columns(2)
        figf = fig_pozo(FUENTE_D)
        cf1.pyplot(figf)
        plt.close(figf)
        cf2.dataframe(pd.DataFrame(NIV_FUENTE, columns=["Elemento", "Cota (msnm)", "Referencia"]).round(2),
                      hide_index=True, width="stretch")
    else:
        c1, c2 = st.columns(2)
        with c1:
            prof_carc = st.number_input("Profundidad del fondo del cárcamo (m bajo terreno)", 1.0, 30.0, 4.0, 0.5)
            ancho = st.number_input("Ancho interior propuesto (m)", 1.0, 30.0, 3.0, 0.5)
            tciclo_def = 5 if Qlps < 30 else (8 if Qlps < 150 else 12)
            tc = st.number_input("Tiempo mínimo de ciclo de arranques (min) [verificar MAPAS]", 2.0, 60.0,
                                 float(tciclo_def), 0.5)
        with c2:
            h_succ = st.number_input("Pérdidas en succión y piezas especiales de la estación (m)", 0.0, 10.0, 0.8, 0.1)
            _man_m = st.checkbox("Capturar manualmente el margen (bordo libre) NAME sobre NAMO", value=False)
            margen = st.number_input("Margen NAME sobre NAMO (m)", 0.0, 2.0, 0.3, 0.1) if _man_m else 0.3
        Ds = D_succ / 1000
        Fr = (Q / (math.pi * Ds ** 2 / 4)) / math.sqrt(G * Ds)
        S = Ds * (1 + 2.3 * Fr)                      # sumergencia mínima (Hydraulic Institute)
        C_auto = max(0.15, 0.3 * Ds)                 # automático: 0.3·D de succión, mín. 0.15 m
        _man_c = st.checkbox(f"Capturar manualmente la separación campana–piso (automática: {C_auto:.2f} m)",
                             value=False)
        C_fondo = st.number_input("Separación campana–piso (m)", 0.05, 2.0, float(round(C_auto, 2)), 0.05) \
            if _man_c else C_auto
        V_util = Q * tc * 60 / 4                     # m3  V = Q·t/4
        z_fondo = z_ground0 - prof_carc
        namino = z_fondo + C_fondo + S
        # área: largo
        # NAMO - NAMINO = V/A
        largo = max(1.5, 1.0)
        A_req_h = 1.0
        # se itera el largo para que la altura útil sea razonable (≈0.8-1.5 m)
        h_util = st.number_input("Altura útil entre NAMINO y NAMO (m)", 0.3, 4.0, 1.0, 0.1)
        A = V_util / h_util
        largo = A / ancho
        namo = namino + h_util
        name_ = namo + margen
        z_low = namino
        z_high = namo
        h_col = h_succ
        st.write(f"Volumen útil V = Q·tc/4 = **{V_util:.1f} m³**; planta **{ancho:.2f} × {largo:.2f} m**")
        st.write(f"Cotas: fondo {z_fondo:.2f} | NAMINO {namino:.2f} | NAMO {namo:.2f} | NAME {name_:.2f}  "
                 f"(sumergencia mínima S = {S:.2f} m)")
        if name_ > z_ground0 - 0.3:
            st.error("El NAME queda muy cerca/sobre el terreno: profundiza el cárcamo o ajusta dimensiones.")
            ALERTAS.append("Cárcamo: NAME demasiado cercano al nivel de terreno.")
        for k, v, u in [("Volumen útil", V_util, "m³"), ("Largo", largo, "m"), ("Ancho", ancho, "m"),
                        ("Sumergencia mínima", S, "m"), ("Cota fondo", z_fondo, "msnm"),
                        ("Cota NAMINO (paro)", namino, "msnm"), ("Cota NAMO (arranque)", namo, "msnm"),
                        ("Cota NAME", name_, "msnm")]:
            add("Cárcamo", k, v, u)
        FUENTE_D = dict(tipo="carcamo", z_t=z_ground0, z_fondo=z_fondo, z_succ=z_fondo + C_fondo,
                        namino=namino, namo=namo, name=name_)
        NIV_FUENTE = [
            ("Terreno en el cárcamo", z_ground0, "Cota de terreno (inicio del perfil)"),
            ("NAME (nivel de aguas máximas extraordinarias)", name_, f"NAMO + {margen:.2f} m"),
            ("NAMO (arranque de bomba)", namo, f"NAMINO + {h_util:.2f} m (volumen útil)"),
            ("NAMINO (paro de bomba)", namino, f"Fondo + holgura + sumergencia = {C_fondo + S:.2f} m"),
            ("Succión de la bomba", z_fondo + C_fondo, f"Fondo + {C_fondo:.2f} m"),
            ("Fondo del cárcamo", z_fondo, f"{prof_carc:.1f} m bajo terreno"),
        ]
        cf1, cf2 = st.columns(2)
        figf = fig_carcamo(FUENTE_D)
        cf1.pyplot(figf)
        plt.close(figf)
        cf2.dataframe(pd.DataFrame(NIV_FUENTE, columns=["Elemento", "Cota (msnm)", "Referencia"]).round(2),
                      hide_index=True, width="stretch")
    add("Bombeo", "Diámetro succión (v≤1.5 m/s)", D_succ, "mm")
    add("Bombeo", "Diámetro descarga bomba (v≤3 m/s)", D_desc, "mm")

# ------------------------------------------------------------------ 4 ACCESORIOS
# puntos altos/bajos
zp_tmp = prof["zp"].to_numpy(float)
hi, lo = puntos_alto_bajo(zp_tmp)
ch_arr = prof["ch"].to_numpy(float)


def sugerir():
    rows = []
    c0, cf = ch_arr[0], ch_arr[-1]
    for el, n in [("Ampliación gradual", 1), ("Válvula de retención (columpio)", 1),
                  ("Válvula de compuerta (abierta)", 1), ("Medidor electromagnético", 1),
                  ("Codo 90°", 1)]:
        rows.append([c0, el, n, K_TABLE[el], np.nan])
    for i in hi:
        rows.append([ch_arr[i], "Válvula de expulsión de aire", 1, 0.0, np.nan])
    for i in lo:
        rows.append([ch_arr[i], "Válvula de desagüe (purga)", 1, 0.0, np.nan])
    rows.append([cf, "Válvula de compuerta (abierta)", 1, K_TABLE["Válvula de compuerta (abierta)"], np.nan])
    rows.append([cf, "Salida a tanque (descarga)", 1, 1.0, np.nan])
    return pd.DataFrame(rows, columns=["Cadenamiento (m)", "Elemento", "Cantidad", "K", "Diámetro (mm)"])


if "acc" not in st.session_state or st.session_state.acc.empty:
    st.session_state.acc = sugerir()

with tabs[3]:
    st.write("Puntos altos detectados (posibles bolsas de aire): " +
             (", ".join(f"{ch_arr[i]:.0f} m" for i in hi) or "ninguno") +
             " | Puntos bajos (desagüe): " + (", ".join(f"{ch_arr[i]:.0f} m" for i in lo) or "ninguno"))
    c1, c2 = st.columns([1, 3])
    with c1:
        if st.button("Regenerar sugerencia automática"):
            st.session_state.acc = sugerir()
            st.session_state.acc_v += 1
            st.rerun()
        long_max = st.number_input("Si hay tramos largos sin punto alto, valvula cada (m)", 100, 3000, 1000, 100)
        st.caption("K en blanco = se toma de la tabla precargada. Para 'Otro (manual)' escribe K.\n"
                   "Diámetro en blanco = diámetro de la línea.")
    with c2:
        acc = st.data_editor(
            st.session_state.acc, num_rows="dynamic", key=f"acc_ed_{st.session_state.acc_v}",
            column_config={
                "Elemento": st.column_config.SelectboxColumn(options=ACC_OPTIONS, required=True),
                "Cadenamiento (m)": st.column_config.NumberColumn(format="%.1f"),
                "Cantidad": st.column_config.NumberColumn(min_value=1, step=1, default=1),
                "K": st.column_config.NumberColumn(format="%.2f"),
                "Diámetro (mm)": st.column_config.NumberColumn(format="%.0f")}, width="stretch")
    st.session_state.acc = acc
    # sin punto alto en tramos largos
    air = sorted(acc.loc[acc["Elemento"] == "Válvula de expulsión de aire", "Cadenamiento (m)"].dropna())
    marks = [ch_arr[0]] + air + [ch_arr[-1]]
    for a_, b_ in zip(marks[:-1], marks[1:]):
        if b_ - a_ > long_max:
            st.warning(f"Tramo {a_:.0f}–{b_:.0f} m sin válvula de aire ({b_ - a_:.0f} m): considera una intermedia.")
    missing = [ch_arr[i] for i in hi if not any(abs(ch_arr[i] - x) < 1 for x in air)]
    for m_ in missing:
        st.warning(f"Punto alto en {m_:.0f} m sin válvula de expulsión de aire.")
        ALERTAS.append(f"Punto alto en cad. {m_:.0f} m sin válvula de expulsión de aire.")
    # diámetros de accesorios
    d_air = commercial_up(max(25, DN / 12), [25, 38, 50, 75, 100, 150, 200])
    d_drain = commercial_up(max(50, DN / 4), [50, 75, 100, 150, 200, 250])
    dacc = pd.DataFrame([
        ["Válvula de compuerta / mariposa / retención en línea", DN, "Igual al diámetro de la línea"],
        ["Válvula de expulsión de aire", d_air, "≈ D/12 (orificio/conexión), mín. 25 mm [verificar]"],
        ["Válvula de desagüe (purga)", d_drain, "≈ D/4, mín. 50 mm [verificar]"],
        ["Succión de bomba", D_succ, "v ≤ 1.5 m/s"],
        ["Descarga de bomba", D_desc, "v ≤ 3.0 m/s (ampliación gradual a DN de línea)"],
    ], columns=["Accesorio", "DN recomendado (mm)", "Criterio"])
    st.write("**Diámetros de accesorios**")
    st.dataframe(dacc, width="stretch")
    SHEETS["Diametros_accesorios"] = dacc
    add("Accesorios", "DN válvula de aire", d_air, "mm")
    add("Accesorios", "DN válvula de desagüe", d_drain, "mm")

# ------------------------------------------------------------------ 5 HIDRÁULICA
# Envejecimiento de la tubería (parámetros editables al final de la pestaña 5; se leen del estado)
_plast = mat_name in ("PVC hidráulico", "PEAD (polietileno AD)", "PRFV (fibra de vidrio)")
ALFA_DEF, DC_DEF = (0.005, 5.0) if _plast else (0.05, 15.0)   # valores de referencia editables (verificar)
age_anios = st.session_state.get("age_anios", 20)
age_alfa = st.session_state.get(f"age_alfa_{mat_name}", ALFA_DEF)
age_dc = st.session_state.get(f"age_dc_{mat_name}", DC_DEF)
eps_env = eps + age_alfa * age_anios                # ε(t) = ε0 + α·t
C_env = Chw * (1 - age_dc / 100.0)                  # C envejecido
cond_dis = st.session_state.get("cond_diseno", "Tubería nueva")
res_new = hidraulica(prof, cover, D, Q, eps, Chw, metodo, acc, z_dest)
res_old = hidraulica(prof, cover, D, Q, eps_env, C_env, metodo, acc, z_dest)
res = res_old if cond_dis.endswith("envejecida") else res_new
TDH_max = res["H0"] - z_low + h_col
TDH_min = res["H0"] - z_high + h_col
pot = potencia(Q, TDH_max, eta_b, eta_m, fs)
dH_surge = a_cel * res["v"] / G
P_op = res["P"].max()
P_st = res["Pst"].max()
P_surge = max(P_op, P_st) + dH_surge
pn_sel = next((p for p in PN_KGCM2 if p * 10 >= P_surge), PN_KGCM2[-1])
hi_r, lo_r = puntos_alto_bajo(res["zp"])

with tabs[4]:
    st.caption(f"Condición de la tubería usada en estos cálculos: **{cond_dis}** "
               "(se cambia en «Datos para la selección de la bomba», al final de esta pestaña).")
    c = st.columns(5)
    c[0].metric("Longitud (m)", f"{res['L']:.1f}")
    c[1].metric("Σ hf fricción (m)", f"{res['hf']:.2f}")
    c[2].metric("Σ hl locales (m)", f"{res['hl']:.2f}")
    c[3].metric("CDT máx (ND/NAMINO) m", f"{TDH_max:.2f}")
    c[4].metric("CDT mín (NE/NAMO) m", f"{TDH_min:.2f}")
    c = st.columns(5)
    c[0].metric("Potencia hidráulica kW", f"{pot['ph']:.2f}")
    c[1].metric("Potencia al eje kW", f"{pot['p_eje']:.2f}")
    c[2].metric("Motor requerido HP", f"{pot['hp_req']:.1f}")
    c[3].metric("Motor comercial HP", f"{pot['hp_com']}")
    c[4].metric("Potencia eléctrica kW", f"{pot['p_elec']:.2f}")
    # presiones
    st.write(f"Presión máx. operación: **{P_op:.1f} m**, estática: **{P_st:.1f} m**, "
             f"sobrepresión estimada (Joukowsky, cierre instantáneo, a={a_cel} m/s): **{dH_surge:.1f} m** → "
             f"Presión máx. de diseño ≈ **{P_surge:.1f} m ≈ {P_surge / 10:.1f} kg/cm²** → clase mínima **{pn_sel} kg/cm²**.")
    st.caption("Estimación conservadora; MAPAS exige análisis de transitorios (golpe de ariete) con tu software.")
    neg = [(res["ch"][i], res["P"][i]) for i in hi_r if res["P"][i] < 2]
    for c_, p_ in neg:
        msg = f"Punto alto en {c_:.0f} m con presión de operación {p_:.1f} m.c.a. (riesgo de presión negativa/bolsa de aire)."
        st.warning(msg)
        ALERTAS.append(msg)
    if res["P"].min() < 0:
        ALERTAS.append("Existen presiones negativas en operación: la línea piezométrica queda bajo la tubería.")
        st.error("Hay presiones negativas en operación: revisa diámetro, posición de bomba o trazo.")
    fg = fig_perfil(prof, res, acc, hi_r, lo_r, pn_sel * 10)
    st.pyplot(fg)
    FIGS.append(("Figura 1. Perfil, línea de energía y presiones.", fg))
    tbl = pd.DataFrame({"Cadenamiento": res["ch"], "Terreno": res["zt"], "Tubería": res["zp"],
                        "Línea de energía": res["LE"].round(2), "Presión operación": res["P"].round(2),
                        "Presión estática": res["Pst"].round(2)})
    st.dataframe(tbl, height=250)
    st.write("Pérdidas locales por accesorio")
    st.dataframe(res["accr"].rename(columns={"c": "Cad. (m)", "n": "Cant.", "d_mm": "Ø (mm)",
                                             "v": "V (m/s)", "hl": "hl (m)"}).round(3))
    SHEETS["Perfil_hidraulico"] = tbl
    SHEETS["Perdidas_locales"] = res["accr"].round(4)
    # costo energético
    tarifa = st.number_input("Tarifa eléctrica ($/kWh) para costo anual estimado", 0.0, 20.0, 2.5, 0.1)
    kwh_anio = pot["p_elec"] * T * 365
    st.write(f"Energía anual ≈ **{kwh_anio:,.0f} kWh**, costo ≈ **${kwh_anio * tarifa:,.0f}**")

    sec = "Hidráulica"
    add("Datos", "Proyecto", proyecto)
    add("Datos", "Fuente", fuente)
    add("Datos", "Equipo", equipo)
    add("Datos", "Caudal de bombeo", Qlps, "l/s")
    add("Datos", "Horas de bombeo", T, "h/día")
    add("Datos", "Destino", destino)
    add(sec, "Material", mat_name)
    add(sec, "Ecuación de fricción", metodo)
    add(sec, "Rugosidad absoluta", eps, "mm")
    add(sec, "Diámetro nominal / interior", f"{DN} / {D * 1000:.1f}", "mm")
    add(sec, "Velocidad", res["v"], "m/s")
    add(sec, "Longitud", res["L"], "m")
    add(sec, "Pérdidas por fricción", res["hf"], "m")
    add(sec, "Pérdidas locales", res["hl"], "m")
    add(sec, "Cota nivel de descarga", z_dest, "msnm")
    add(sec, "Cota nivel mínimo fuente (ND/NAMINO)", z_low, "msnm")
    add(sec, "Carga dinámica total máxima", TDH_max, "m")
    add(sec, "Carga dinámica total mínima", TDH_min, "m")
    add(sec, "Presión máxima de operación", P_op, "m.c.a.")
    add(sec, "Sobrepresión estimada", dH_surge, "m")
    add(sec, "Clase mínima de tubería", pn_sel, "kg/cm²")
    add("Bomba", "Potencia hidráulica", pot["ph"], "kW")
    add("Bomba", "Eficiencia bomba / motor", f"{eta_b} / {eta_m}")
    add("Bomba", "Potencia al eje", pot["p_eje"], "kW")
    add("Bomba", "Motor requerido (FS)", pot["hp_req"], "HP")
    add("Bomba", "Motor comercial", pot["hp_com"], "HP")
    add("Bomba", "Potencia eléctrica", pot["p_elec"], "kW")

# ------------------------------------------------------------------ 7 TANQUES
with tabs[6]:
    st.write(f"Qmd (definido en la pestaña Datos; con él se diseña el tanque de regulación) = **{Qmd:.2f} l/s**")
    met = st.radio("Método de regulación", ["Coeficiente de regulación (MAPAS, 24 h)", "Curva masa (patrón editable)"])
    if T < 24 and met.startswith("Coef"):
        st.warning("El coeficiente 14.58 corresponde a bombeo continuo 24 h; para menos horas usa curva masa.")
    cum = None
    if met.startswith("Coef"):
        Cr = st.number_input("Coeficiente de regulación Cr (m³ por l/s)", 1.0, 40.0, CR_24H, 0.01)
    else:
        patron = st.data_editor(pd.DataFrame({"Hora": range(24), "Coef. demanda horaria": PATRON_DEMANDA}),
                                key="patron", width="stretch")
        Cr, cum = cr_curva_masa(patron["Coef. demanda horaria"].tolist(), T, int(h0))
        st.write(f"Cr calculado por curva masa = **{Cr:.2f} m³/(l/s)** (referencia 24 h MAPAS: {CR_24H})")
        figm = fig_curva_masa(cum)
        st.pyplot(figm)
        FIGS.append(("Figura 2. Curva masa de regulación.", figm))
    Vreg = Cr * Qmd
    c1, c2 = st.columns(2)
    with c1:
        t_res = st.number_input("Tiempo de almacenamiento de reserva/emergencia (h de Qmd)", 0.0, 72.0, 0.0, 1.0)
        Vres = Qmd * 3.6 * t_res
    with c2:
        h_tq = st.number_input("Tirante útil propuesto (m)", 1.0, 15.0,
                               float(min(max(h_llegada, 1.0), 15.0)) if cist_recibe_piso
                               else (float(min(max(d_tir, 1.0), 15.0)) if desg else 3.5), 0.1)
        forma = st.radio("Forma", ["Circular", "Rectangular (largo = 1.5 × ancho)"], horizontal=True)
    pct_el = 100.0
    if destino == "Cisterna + tanque elevado":
        pct_el = st.slider("% del volumen de regulación alojado en el tanque elevado", 0, 100, 40)
    pct_el_n, pct_pi = 0, 0
    if destino == "Cisterna + tanque de piso":
        pct_pi = st.slider("% del volumen de regulación alojado en el tanque de piso", 0, 100, 60, key="pct_pi")
    elif destino == "Tanque de piso + tanque elevado":
        pct_el_n = st.slider("% del volumen de regulación alojado en el tanque elevado", 0, 100, 40, key="pct_el_n")
    elif destino == "Cisterna + tanque de piso + tanque elevado":
        _s1, _s2 = st.columns(2)
        pct_el_n = _s1.slider("% de la regulación en el tanque elevado", 0, 100, 40, key="pct_el_3")
        pct_pi = _s2.slider("% de la regulación en el tanque de piso", 0, 100, 60, key="pct_pi_3")
        if pct_el_n + pct_pi > 100:
            st.error("La suma de porcentajes supera 100 %: se ajusta el del tanque de piso.")
            pct_pi = 100 - pct_el_n
        st.caption(f"La cisterna (almacenamiento) recibe el {100 - pct_el_n - pct_pi} % restante de la regulación "
                   "más toda la reserva.")
    add("Tanques", "Qmd", Qmd, "l/s")
    add("Tanques", "Coeficiente de regulación", Cr, "m³/(l/s)")
    add("Tanques", "Volumen de regulación", Vreg, "m³")
    add("Tanques", "Volumen de reserva", Vres, "m³")

    def dims(V, h, n=1):
        """Dimensiones de UNA unidad (volumen V/n y tirante h)."""
        if V <= 0 or h <= 0:
            return "sin volumen (0 m³)"
        A = (V / n) / h
        if forma == "Circular":
            txt = f"Ø {math.sqrt(4 * A / math.pi):.2f} m × {h:.2f} m"
        else:
            b_ = math.sqrt(A / 1.5)
            txt = f"{1.5 * b_:.2f} × {b_:.2f} m × {h:.2f} m"
        return f"{n} × ({txt})" if n > 1 else txt

    tq = []
    REG_K = None
    if destino == "Cisterna + tanque elevado":
        V_el = Vreg * pct_el / 100
        V_ci = Vreg - V_el + Vres
        tq = [("Tanque elevado", V_el), ("Cisterna", V_ci)]
    elif destino == "Tanque elevado":
        tq = [("Tanque elevado", Vreg + Vres)]
    elif destino == "Cisterna":
        tq = [("Cisterna", Vreg + Vres)]
    elif destino == "Cisterna + tanque de piso":
        V_pi = Vreg * pct_pi / 100
        tq = [("Tanque de regulación", V_pi), ("Cisterna", Vreg - V_pi + Vres)]
        REG_K = {"Tanque de regulación": V_pi, "Cisterna": Vreg - V_pi}
    elif destino == "Tanque de piso + tanque elevado":
        V_el = Vreg * pct_el_n / 100
        tq = [("Tanque elevado", V_el), ("Tanque de regulación", Vreg - V_el + Vres)]
        REG_K = {"Tanque elevado": V_el, "Tanque de regulación": Vreg - V_el}
    elif destino == "Cisterna + tanque de piso + tanque elevado":
        V_el, V_pi = Vreg * pct_el_n / 100, Vreg * pct_pi / 100
        V_ci_r = Vreg - V_el - V_pi
        tq = [("Tanque elevado", V_el), ("Tanque de regulación", V_pi), ("Cisterna", V_ci_r + Vres)]
        REG_K = {"Tanque elevado": V_el, "Tanque de regulación": V_pi, "Cisterna": V_ci_r}
        if V_ci_r + Vres <= 0:
            st.info("La cisterna de almacenamiento queda en 0 m³: captura horas de reserva (arriba) o asígnale "
                    "parte de la regulación.")
    else:
        tq = [("Tanque de regulación", Vreg + Vres)]

    # ---------------- división en varias unidades (opcional)
    st.markdown("#### División en varias unidades (opcional)")
    V_ref = st.number_input("Volumen de referencia por unidad (m³) solo para sugerir la división "
                            "[criterio del proyectista]", 50.0, 50000.0, 1000.0, 50.0)
    N_UN = {}
    for (n_, v_), col_ in zip(tq, st.columns(len(tq))):
        sug = max(1, math.ceil(v_ / V_ref))
        N_UN[n_] = int(col_.number_input(f"Unidades de: {n_}", 1, 20, 1, 1, key=f"nu_{n_}"))
        col_.caption(f"Volumen total {v_:.1f} m³ → " +
                     (f"se sugiere dividir en {sug} unidades" if sug > 1 else "una unidad es suficiente"))

    # ---------------- igualar planta de cisterna y tanque (opcional)
    H_U = {n_: h_tq for n_, _ in tq}
    igualar = False
    if len(tq) >= 2:
        st.markdown("#### Igualar la planta de cisterna y tanque (opcional)")
        igualar = st.checkbox("Igualar la sección en planta (para colocar la cisterna bajo el tanque o hacer "
                              "una sola estructura). Por defecto quedan separados.", value=False)
        if igualar:
            grp = [n_ for n_, _ in tq]
            if len(grp) > 2:
                grp = st.multiselect("Estructuras cuya planta se iguala", grp, default=grp, key="grp_ig")
            tq_g = [(n_, v_) for n_, v_ in tq if n_ in grp]
            if len(tq_g) < 2:
                st.warning("Selecciona al menos dos estructuras; se igualan todas.")
                tq_g = list(tq)
            if len({N_UN[n_] for n_, _ in tq_g}) > 1:
                st.warning("Las estructuras tienen distinto número de unidades; para apilarlas usa el mismo número.")
            modo_ig = st.radio("Planta común", ["Tomar la de la estructura de mayor volumen",
                                                 "Capturar diámetro (circular) o ancho (rectangular)"],
                               horizontal=True)
            A0 = max(max((v_ / N_UN[n_]) / h_tq for n_, v_ in tq_g), 1.0)
            if modo_ig.startswith("Tomar"):
                A_c = A0
            elif forma == "Circular":
                d_c = st.number_input("Diámetro común (m)", 1.0, 80.0,
                                      float(min(max(round(math.sqrt(4 * A0 / math.pi), 1), 1.0), 80.0)), 0.1)
                A_c = math.pi * d_c ** 2 / 4
            else:
                b_c = st.number_input("Ancho común (m); largo = 1.5 × ancho", 1.0, 60.0,
                                      float(min(max(round(math.sqrt(A0 / 1.5), 1), 1.0), 60.0)), 0.1)
                A_c = 1.5 * b_c ** 2
            for n_, v_ in tq_g:
                H_U[n_] = (v_ / N_UN[n_]) / A_c if v_ > 0 else h_tq
            if any(v_ <= 0 for _, v_ in tq_g):
                st.info("Una de las estructuras tiene volumen 0 m³ (revisa el % en el tanque elevado o la reserva): "
                        "no se le ajusta el tirante.")
            st.write("Tirantes resultantes con la planta común: " +
                     "; ".join(f"**{n_}**: {H_U[n_]:.2f} m" for n_, _ in tq_g) +
                     f". Planta por unidad: {A_c:.1f} m².")
            if any((H_U[n_] < 1.0 or H_U[n_] > 10.0) for n_, v_ in tq_g if v_ > 0):
                st.warning("Algún tirante queda fuera de 1.0–10.0 m (referencia); revisa la viabilidad constructiva.")

    def _contenido(n_, v_):
        if REG_K is None:
            return "Reg. + reserva" if n_ != "Tanque elevado" else "Regulación"
        r_ = REG_K.get(n_, 0.0)
        e_ = max(v_ - r_, 0.0)
        return "Reg. + reserva" if (r_ > 0.05 and e_ > 0.05) else ("Reserva" if e_ > 0.05 else "Regulación")

    rows_t = []
    for n_, v_ in tq:
        n_u, h_i = N_UN[n_], H_U[n_]
        rows_t.append([n_, round(v_, 1), n_u, round(v_ / n_u, 1), dims(v_, h_i, n_u),
                       _contenido(n_, v_)])
        add("Tanques", f"{n_} - volumen", v_, "m³")
        add("Tanques", f"{n_} - dimensiones", dims(v_, h_i, n_u))
        if n_u > 1:
            add("Tanques", f"{n_} - unidades", n_u)
            add("Tanques", f"{n_} - volumen por unidad", v_ / n_u, "m³")
    tdf = pd.DataFrame(rows_t, columns=["Estructura", "Volumen (m³)", "Unidades", "Volumen por unidad (m³)",
                                        "Dimensiones", "Contenido"])
    st.dataframe(tdf, width="stretch")
    _it = [(n_, N_UN[n_], (v_ / N_UN[n_]) / H_U[n_]) for n_, v_ in tq if v_ > 0 and H_U[n_] > 0]
    if _it:
        fig_pl = fig_planta(_it, forma)
        st.pyplot(fig_pl)
        FIGS.append(("Figura. Planta de las estructuras de almacenamiento.", fig_pl))
    if igualar:
        st.success("Plantas iguales: la cisterna puede alojarse bajo el tanque o formar una sola estructura.")
    st.info("Ajusta volúmenes comerciales/constructivos. Reserva contra incendio y emergencia: "
            "definir según lineamientos del organismo operador.")
    SHEETS["Tanques"] = tdf

    # ---------------- niveles de las estructuras
    st.markdown("#### Niveles de las estructuras de almacenamiento")
    _man_b = st.checkbox("Capturar manualmente el bordo libre (por defecto 0.50 m)", value=False)
    bordo = st.number_input("Bordo libre (m)", 0.1, 3.0, 0.5, 0.05) if _man_b else 0.5

    def _key(nombre):
        return "Tanque elevado" if "elevado" in nombre else ("Cisterna" if "Cisterna" in nombre
                                                              else "Tanque de regulación")

    def estr_(nombre, namo, z_t, recibe=False):
        k_ = _key(nombre)
        h_i = H_U.get(k_, h_tq)
        return dict(nombre=nombre, z_t=z_t, fondo=namo - h_i, namin=namo - h_i, namo=namo,
                    corona=namo + bordo, recibe=recibe, h=h_i, n=N_UN.get(k_, 1))

    hC, hE = H_U.get("Cisterna", h_tq), H_U.get("Tanque elevado", h_tq)

    def _aplicar_h(v):
        st.session_state["h_llegada_in"] = min(max(float(v), 0.0), 80.0)

    base_man, base_tipo, base_h = False, None, 0.0
    hay_piso = destino == "Tanque de regulación superficial (piso)" or (
        destino in DEST_NUEVOS and "Tanque de piso" in DEST_ESTR[destino])
    if hay_piso:
        st.markdown("#### Base del tanque de piso (opcional)")
        base_man = st.checkbox("Proponer manualmente la altura de la base (firme de concreto o estructura de acero)",
                               value=False, key="base_on")
        if base_man:
            bb1, bb2 = st.columns(2)
            base_tipo = bb1.selectbox("Tipo de base", ["Firme de concreto", "Estructura de acero"], key="base_tipo")
            base_h = bb2.number_input("Altura de la base sobre el terreno (m) [valor de referencia, ajústalo]",
                                      0.0, 10.0, 0.30 if base_tipo == "Firme de concreto" else 1.50, 0.05,
                                      key=f"base_h_{base_tipo}")
            add("Tanques", "Base del tanque", f"{base_tipo}, {base_h:.2f} m sobre el terreno")

    ESTR = []
    if destino in DEST_NUEVOS:
        _nom = {"Cisterna": "Cisterna (sobre terreno, tipo piso)" if cist_piso else "Cisterna",
                "Tanque de piso": "Tanque de regulación (piso)", "Tanque elevado": "Tanque elevado"}
        for k_ in [llega_en] + [o_ for o_ in DEST_ESTR[destino] if o_ != llega_en]:
            nm_k = _nom[k_]
            h_k = H_U.get(_key(nm_k), h_tq)
            if k_ == llega_en:
                if k_ == "Tanque de piso" and base_man:
                    ESTR.append(estr_(nm_k, z_end + base_h + h_k, z_end, True))
                    h_req = base_h + h_k + caida
                    if abs(h_req - h_llegada) > 0.01:
                        st.warning(f"La hidráulica usa h = {h_llegada:.2f} m (pestaña Datos), pero base + tirante = "
                                   f"{h_req:.2f} m. Para que la carga de bombeo corresponda a esta geometría, aplica el valor.")
                        if desg:
                            st.info("El desglose de h está activo en Datos: ajusta allí la base/torre y el tirante.")
                        else:
                            st.button(f"Aplicar h = {h_req:.2f} m en Datos", on_click=_aplicar_h,
                                      args=(round(h_req, 2),), key="btn_apply_h_n")
                    else:
                        st.success("La altura de descarga h de Datos coincide con base + tirante.")
                else:
                    ESTR.append(estr_(nm_k, z_dest - caida, z_end, True))
            elif k_ == "Cisterna":
                if cist_piso:
                    ESTR.append(estr_(nm_k, z_end + h_k, z_end))
                else:
                    prof_ci = st.number_input("Profundidad del fondo de la cisterna bajo el terreno (m)", 0.0, 15.0,
                                              float(min(h_k, 15.0)), 0.5, key="prof_ci_n")
                    ESTR.append(estr_(nm_k, z_end - prof_ci + h_k, z_end))
            elif k_ == "Tanque de piso":
                ESTR.append(estr_(nm_k, z_end + (base_h if base_man else 0.0) + h_k, z_end))
            else:
                torre = st.number_input("Altura del fondo del tanque elevado sobre el terreno (m)", 0.0, 60.0, 12.0,
                                        1.0, key="torre_n")
                ESTR.append(estr_(nm_k, z_end + torre + h_k, z_end))
            if k_ == "Tanque de piso" and base_man:
                ESTR[-1]["base_tipo"], ESTR[-1]["base_h"] = base_tipo, base_h
        st.caption("La línea descarga en la primera estructura. Las demás se ubican en el mismo sitio: si quedan más "
                   "bajas se llenan por gravedad; si quedan más altas requieren rebombeo (no se calcula aquí).")
    elif destino != "Cisterna + tanque elevado":
        nm = {"Tanque de regulación superficial (piso)": "Tanque de regulación (piso)",
              "Tanque elevado": "Tanque elevado", "Cisterna": "Cisterna"}[destino]
        if base_man:
            ESTR.append(estr_(nm, z_end + base_h + H_U.get(_key(nm), h_tq), z_end, True))
            ESTR[-1]["base_tipo"], ESTR[-1]["base_h"] = base_tipo, base_h
            h_req = base_h + ESTR[-1]["h"] + caida
            if abs(h_req - h_llegada) > 0.01:
                st.warning(f"La hidráulica usa h = {h_llegada:.2f} m (pestaña Datos), pero base + tirante = "
                           f"{h_req:.2f} m. Para que la carga de bombeo corresponda a esta geometría, aplica el valor.")
                if desg:
                    st.info("El desglose de h está activo en Datos: ajusta allí la base/torre y el tirante.")
                else:
                    st.button(f"Aplicar h = {h_req:.2f} m en Datos", on_click=_aplicar_h, args=(round(h_req, 2),),
                              key="btn_apply_h")
            else:
                st.success("La altura de descarga h de Datos coincide con base + tirante.")
        else:
            ESTR.append(estr_(nm, z_dest - caida, z_end, True))
    else:
        ESTR.append(estr_("Cisterna (sobre terreno, tipo piso)" if cist_recibe_piso else llega_en,
                          z_dest - caida, z_end, True))
        otro = "Cisterna" if llega_en == "Tanque elevado" else "Tanque elevado"
        if otro == "Cisterna":
            if cist_piso:
                st.caption("Cisterna sobre el terreno (tipo piso): fondo a nivel de terreno, sin excavación.")
                ESTR.append(estr_("Cisterna (sobre terreno, tipo piso)", z_end + hC, z_end))
            else:
                prof_ci = st.number_input("Profundidad del fondo de la cisterna bajo el terreno (m)", 0.0, 15.0,
                                          float(min(hC, 15.0)), 0.5)
                ESTR.append(estr_("Cisterna", z_end - prof_ci + hC, z_end))
        else:
            torre = st.number_input("Altura del fondo del tanque elevado sobre el terreno (m)", 0.0, 60.0, 12.0, 1.0)
            ESTR.append(estr_("Tanque elevado", z_end + torre + hE, z_end))
        st.caption("La línea descarga en la primera estructura; la segunda se ubica en el mismo sitio. "
                   "El rebombeo entre ambas no se calcula en esta app.")
    if cist_recibe_piso and abs(ESTR[0]["fondo"] - z_end) > 0.05:
        st.warning(f"Cisterna sobre el terreno (sin excavación): el fondo debe quedar a nivel de terreno. "
                   f"Ajusta h en Datos ≈ tirante útil ({ESTR[0]['h']:.2f} m) o el tirante ≈ h ({h_llegada:.2f} m).")
    cn1, cn2 = st.columns([1.2, 1])
    fig_t = fig_tanques(ESTR)
    cn1.pyplot(fig_t)
    FIGS.append(("Figura. Niveles de las estructuras de almacenamiento.", fig_t))
    ndf = pd.DataFrame([[e["nombre"], e["fondo"], e["namo"], e["corona"], e["fondo"] - e["z_t"], e["n"], e["h"]]
                        for e in ESTR],
                       columns=["Estructura", "Fondo = NAMIN (msnm)", "NAMO (msnm)", "Corona (msnm)",
                                "Fondo sobre(+)/bajo(−) terreno (m)", "Unidades", "Tirante útil (m)"]).round(2)
    cn2.dataframe(ndf, hide_index=True, width="stretch")
    for e in ESTR:
        rel = e["fondo"] - e["z_t"]
        if rel > 0.05:
            cn2.caption(f"{e['nombre']}: fondo a {rel:.1f} m sobre el terreno (torre o base).")
        elif rel < -0.05:
            cn2.caption(f"{e['nombre']}: semienterrado, fondo {abs(rel):.1f} m bajo el terreno.")
    SHEETS["Niveles_tanques"] = ndf
    TABLAS_DOC.append(("Volúmenes y dimensiones de las estructuras de almacenamiento", tdf))
    TABLAS_DOC.append(("Niveles de las estructuras de almacenamiento", ndf))
    _vol = dict(tq)
    for e in ESTR:
        _k = "Tanque elevado" if "elevado" in e["nombre"] else ("Cisterna" if "Cisterna" in e["nombre"]
                                                                  else "Tanque de regulación")
        _v = _vol.get(_k, 0.0)
        _rel = e["fondo"] - e["z_t"]
        _pos = (f"el fondo queda a {_rel:.2f} m sobre el terreno ({e.get('base_tipo', 'torre o base').lower()})" if _rel > 0.05 else
                f"el fondo queda a {abs(_rel):.2f} m bajo el terreno (semienterrado)" if _rel < -0.05 else
                "el fondo queda a nivel de terreno (sin excavación)")
        PARRAFOS_DOC.append(f"{e['nombre']}: volumen {_v:.1f} m³"
                            + (f" ({e['n']} unidades de {_v / e['n']:.1f} m³ c/u)" if e["n"] > 1 else "")
                            + f"; dimensiones {dims(_v, e['h'], e['n'])}; "
                            f"cota de fondo {e['fondo']:.2f} msnm, NAMO {e['namo']:.2f} msnm, corona "
                            f"{e['corona']:.2f} msnm; {_pos}.")

    # ---------------- carga de bombeo según la estructura que recibe la línea
    st.markdown("#### Carga de bombeo según la estructura que recibe la línea")
    with st.expander("¿Cómo se calcula la altura de llegada y la línea de energía (LE) de la bomba?"):
        st.markdown(
            "1. **Cota de descarga** `z_dest = terreno final + h`. La altura `h` es el nivel máximo del agua (NAMO) "
            "sobre el terreno final: **piso** = base + tirante; **elevado** = torre + tirante; **cisterna** = fondo "
            "(negativo si está enterrada) + tirante. Si la boca de descarga queda **sobre** el agua (caída libre), se suma "
            "esa caída; si llega **sumergida**, no.\n"
            "2. **LE al final** `= z_dest + K·v²/2g` (K = 1.0 de salida a tanque).\n"
            "3. **Hacia aguas arriba:** `LE(x) = z_dest + Σhf aguas abajo + Σhl aguas abajo`.\n"
            "4. **CDT** `= LE en la descarga de la bomba − nivel de la fuente (ND o NAMINO) + pérdidas en columna/succión`.\n\n"
            "Como las pérdidas no dependen de `z_dest`, cambiar de estructura receptora mueve la CDT exactamente "
            "lo que cambia su cota de descarga: un tanque elevado exige más carga que uno de piso.")
    rec = ESTR[0]
    perd = res["hf"] + res["hl"]
    rows_c = []
    for e in ESTR:
        z_k = e["namo"] + caida
        cdt_k = perd + z_k - z_low + h_col
        if e is rec:
            alim = "Recibe la línea"
        elif e["namo"] > rec["namo"] + 0.05:
            alim = "Más alta que la receptora: requiere rebombeo (no calculado)"
        else:
            alim = "Más baja: se llena por gravedad desde la receptora"
        rows_c.append([e["nombre"], round(e["namo"], 2), round(z_k, 2), round(z_k - z_end, 2), round(cdt_k, 2),
                       round(cdt_k - TDH_max, 2), alim])
    dfr = pd.DataFrame(rows_c, columns=["Estructura", "NAMO (msnm)", "Cota de descarga (msnm)",
                                        "h sobre terreno final (m)", "CDT máx. si la línea llega aquí (m)",
                                        "Diferencia vs. CDT de diseño (m)", "Alimentación entre estructuras"])
    st.dataframe(dfr, hide_index=True, width="stretch")
    st.caption("La CDT de diseño corresponde a la primera fila (estructura que recibe la línea). Una diferencia distinta "
               "de cero en esa fila indica que la h de Datos no coincide con la geometría de la estructura.")
    SHEETS["Carga_segun_estructura"] = dfr
    TABLAS_DOC.append(("Carga de bombeo según la estructura que recibe la línea",
                       dfr[["Estructura", "Cota de descarga (msnm)", "CDT máx. si la línea llega aquí (m)",
                            "Alimentación entre estructuras"]]))
    PARRAFOS_DOC.append(f"La línea de conducción descarga en: {rec['nombre']} "
                        f"({'descarga libre sobre el NAMO' if caida > 0 else 'llegada al nivel del agua'}). Cota de "
                        f"descarga {z_dest:.2f} msnm = terreno final {z_end:.2f} + h {h_llegada:.2f} m; "
                        f"carga dinámica total máxima {TDH_max:.2f} m.")

    # ---------------- alimentación a la red
    st.markdown("#### Alimentación a la red (se diseña con Qmh)")
    st.write(f"Qmh = **{Qmh:.2f} l/s**. La línea de alimentación y la red quedan fuera del alcance de esta "
             "memoria; se da un diámetro orientativo.")
    dq = pd.DataFrame([[v_, round(math.sqrt(4 * Qmh / 1000 / (math.pi * v_)) * 1000),
                        commercial_up(math.sqrt(4 * Qmh / 1000 / (math.pi * v_)) * 1000)] for v_ in (1.0, 1.5, 2.0)],
                      columns=["Velocidad (m/s)", "Ø teórico (mm)", "DN comercial ≥ (mm)"])
    st.dataframe(dq, hide_index=True)
    add("Tanques", "Qmh (alimentación a red)", float(Qmh), "l/s")

# ------------------------------------------------------------------ 5b ESQUEMA GENERAL Y NIVELES
with tabs[4]:
    st.markdown("---")
    st.subheader("Esquema general de niveles (no a escala; las cotas rotuladas son las calculadas)")
    fsys = fig_sistema(FUENTE_D, res, ESTR)
    st.pyplot(fsys)
    FIGS.append(("Figura 3. Esquema general de niveles: fuente, línea de conducción y destino.", fsys))

    with st.expander("📥 Esquema ejecutivo para presentaciones (alta calidad: PNG, SVG, PDF)"):
        st.caption("Versión rediseñada de la Figura 3: texto grande, encabezados por bloque y franja de datos clave. "
                   "Para PowerPoint: SVG (vectorial, editable) o PNG de 300 dpi; para Keynote / Google Slides: PDF o PNG.")
        o1, o2, o3, o4 = st.columns(4)
        tam = o1.selectbox("Tamaño", ["Diapositiva 16:9 (13.33 × 7.5 in)", "Diapositiva 4:3 (10 × 7.5 in)",
                                      "A4 horizontal (11.69 × 8.27 in)"], key="exe_size")
        fesc = o2.slider("Tamaño de texto", 0.8, 1.4, 1.0, 0.05, key="exe_fs")
        dpi_png = o3.selectbox("Resolución PNG (dpi)", [150, 200, 300], index=2, key="exe_dpi")
        fondo = o4.selectbox("Fondo", ["Blanco", "Transparente"], key="exe_bg")
        x1, x2, x3 = st.columns(3)
        k_tit = x1.checkbox("Incluir título", True, key="exe_t")
        k_kpi = x2.checkbox("Incluir franja de datos clave", True, key="exe_k")
        svg_txt = x3.checkbox("SVG con texto editable (requiere la fuente instalada)", False, key="exe_svgtxt")
        y1, y2, _y3 = st.columns(3)
        k_alto = y1.checkbox("Mostrar cuadro «Punto más alto» y su línea", False, key="exe_alto")
        k_lleg = y2.checkbox("Mostrar cuadro «Llegada» y su línea", False, key="exe_lleg")
        size_exe = {"Diapositiva 16:9 (13.33 × 7.5 in)": (13.333, 7.5), "Diapositiva 4:3 (10 × 7.5 in)": (10.0, 7.5),
                    "A4 horizontal (11.69 × 8.27 in)": (11.69, 8.27)}[tam]
        info_exe = dict(
            titulo=proyecto, fuente=fuente,
            subtitulo="Esquema general de niveles · no a escala · cotas en msnm",
            pie=f"Memoria de línea de bombeo · {dt.date.today():%d/%m/%Y}",
            callout=f"CDT máx. {TDH_max:.1f} m\nhf {res['hf']:.1f} m · hl {res['hl']:.1f} m\nv = {res['v']:.2f} m/s",
            kpis=[("Caudal de bombeo", f"{Qlps:.1f} l/s"), ("Carga dinámica total", f"{TDH_max:.1f} m"),
                  ("Motor", f"{pot['hp_com']} HP"), (f"Tubería · {mat_name}", f"Ø{DN} mm · {res['L'] / 1000:.2f} km"),
                  ("Presión máx. de diseño", f"{P_surge / 10:.1f} kg/cm²"),
                  ("Almacenamiento", f"{tdf['Volumen (m³)'].sum():.0f} m³")])
        fex = fig_sistema_exec(FUENTE_D, res, ESTR, info_exe, size=size_exe, fscale=fesc, kpi=k_kpi, titulo=k_tit,
                                 show_alto=k_alto, show_llegada=k_lleg)
        _prev = io.BytesIO()
        fex.savefig(_prev, format="png", dpi=110)
        st.image(_prev)
        if st.button("Preparar archivos de alta calidad (PNG / SVG / PDF)", key="exe_prep"):
            with st.spinner("Generando archivos…"):
                files_exe = {}
                with plt.rc_context({"svg.fonttype": "none" if svg_txt else "path", "pdf.fonttype": 42}):
                    for ext_, kw_ in [("png", dict(dpi=dpi_png)), ("svg", {}), ("pdf", {})]:
                        bio_ = io.BytesIO()
                        fex.savefig(bio_, format=ext_, transparent=(fondo == "Transparente"), **kw_)
                        files_exe[ext_] = bio_.getvalue()
                st.session_state["exe_files"] = files_exe
        plt.close(fex)
        if "exe_files" in st.session_state:
            slug_ = re.sub(r"[^A-Za-z0-9_-]+", "_", proyecto).strip("_") or "esquema"
            d1, d2, d3 = st.columns(3)
            d1.download_button("Descargar PNG", st.session_state["exe_files"]["png"],
                               f"esquema_niveles_{slug_}.png", "image/png", key="dl_exe_png")
            d2.download_button("Descargar SVG (vectorial)", st.session_state["exe_files"]["svg"],
                               f"esquema_niveles_{slug_}.svg", "image/svg+xml", key="dl_exe_svg")
            d3.download_button("Descargar PDF (vectorial)", st.session_state["exe_files"]["pdf"],
                               f"esquema_niveles_{slug_}.pdf", "application/pdf", key="dl_exe_pdf")
            st.caption("Si cambiaste datos u opciones, vuelve a pulsar «Preparar archivos».")
    i_hi = int(np.argmax(res["zp"]))
    i_lo = int(np.argmin(res["zp"]))
    NIV = [("Fuente", e_, c_, r_) for e_, c_, r_ in NIV_FUENTE]
    NIV += [
        ("Bombeo", "Cabezal de descarga (inicio de la tubería)", res["zp"][0], "Cota de tubería al inicio del perfil"),
        ("Bombeo", "Energía entregada por la bomba (nivel mín. fuente + CDT máx.)", z_low + TDH_max,
         f"{z_low:.2f} + {TDH_max:.2f} m"),
        ("Bombeo", "Línea de energía a la descarga (tras columna/succión)", res["H0"], "Inicio de la línea de conducción"),
        ("Línea", "Punto más alto de la tubería", res["zp"][i_hi], f"Cad. {res['ch'][i_hi]:.0f} m"),
        ("Línea", "Línea de energía en el punto más alto", res["LE"][i_hi],
         f"Presión de operación {res['P'][i_hi]:.1f} m.c.a."),
        ("Línea", "Punto más bajo de la tubería", res["zp"][i_lo], f"Cad. {res['ch'][i_lo]:.0f} m"),
        ("Línea", "Fin de la tubería (llegada)", res["zp"][-1], f"Cad. {res['ch'][-1]:.0f} m"),
        ("Destino", "Terreno en el punto final", z_end, "Último punto del perfil"),
        ("Destino", "Nivel máximo de descarga (NAMO de la estructura receptora)", z_dest - caida,
         f"Terreno + h = {h_llegada - caida:.1f} m"),
        *([("Destino", "Cota de la boca de descarga (NAMO + caída libre)", z_dest, f"Caída libre {caida:.2f} m")]
          if caida > 0 else []),
    ]
    for e in ESTR:
        NIV += [("Destino", f"{e['nombre']}: fondo = NAMIN", e["fondo"], f"{e['fondo'] - e['z_t']:+.2f} m respecto al terreno"),
                ("Destino", f"{e['nombre']}: NAMO", e["namo"], f"Tirante útil {e['h']:.2f} m" + (f"; {e['n']} unidades" if e["n"] > 1 else "")),
                ("Destino", f"{e['nombre']}: corona", e["corona"], f"NAMO + bordo libre {bordo:.2f} m")]
    nivdf = pd.DataFrame(NIV, columns=["Grupo", "Elemento", "Cota (msnm)", "Referencia"])
    nivdf["Cota (msnm)"] = nivdf["Cota (msnm)"].astype(float).round(2)
    st.subheader("Tabla de niveles de la instalación")
    st.dataframe(nivdf, hide_index=True, width="stretch")
    SHEETS["Niveles"] = nivdf
    for _, r_ in nivdf.iterrows():
        add("Niveles: " + r_["Grupo"], r_["Elemento"], float(r_["Cota (msnm)"]), "msnm")

# ------------------------------------------------------------------ 5c DATOS PARA LA SELECCIÓN DE LA BOMBA
HV_T = [5, 10, 15, 20, 25, 30, 35, 40, 50]
HV_M = [0.09, 0.12, 0.17, 0.24, 0.32, 0.43, 0.57, 0.75, 1.26]     # presión de vapor del agua (m.c.a.)

with tabs[4]:
    st.markdown("---")
    st.subheader("Datos para la selección de la bomba")

    # ---- 1. tubería nueva vs envejecida
    st.markdown("##### 1. Tubería nueva vs. envejecida")
    st.caption("Los cálculos base usan tubería nueva. Aquí se verifica la envejecida con ε(t) = ε₀ + α·t "
               "(Darcy-Weisbach) o con C reducido (Hazen-Williams). Los valores por defecto son de referencia: "
               "ajústalos con el criterio de MAPAS / calidad del agua.")
    a1, a2, a3, a4 = st.columns(4)
    a1.number_input("Años de servicio considerados", 0, 60, 20, 1, key="age_anios")
    a2.number_input("Incremento anual de rugosidad α (mm/año)", 0.0, 1.0, ALFA_DEF, 0.01,
                    key=f"age_alfa_{mat_name}", format="%.3f")
    a3.number_input("Reducción de C Hazen-Williams (%)", 0.0, 50.0, DC_DEF, 1.0, key=f"age_dc_{mat_name}")
    a4.radio("Dimensionar la bomba con:", ["Tubería nueva", "Tubería envejecida"], key="cond_diseno")
    filas = []
    for etq, rr, e_, c_ in [("Nueva", res_new, eps, Chw), ("Envejecida", res_old, eps_env, C_env)]:
        for niv, zs in [("Nivel mínimo (ND / NAMINO)", z_low), ("Nivel máximo (NE / NAMO)", z_high)]:
            H_ = rr["H0"] - zs + h_col
            p_ = potencia(Q, H_, eta_b, eta_m, fs)
            filas.append([etq, niv, round(e_, 4), round(c_, 1), round(rr["hf"], 2), round(H_, 2),
                          round(p_["p_eje"], 2), round(p_["p_eje"] / 0.7457, 1), round(p_["hp_req"], 1),
                          p_["hp_com"], round(p_["p_elec"], 2)])
    if metodo.startswith("Darcy") and age_alfa * age_anios <= 0:
        st.info("Con α = 0 o 0 años la rugosidad envejecida es igual a la nueva (Darcy-Weisbach): no habrá diferencia.")
    dfc = pd.DataFrame(filas, columns=["Tubería", "Nivel de la fuente", "ε (mm)", "C (H-W)", "hf línea (m)",
                                       "CDT (m)", "Potencia al freno BHP (kW)", "BHP (HP)",
                                       "Motor requerido (HP)", "Motor comercial (HP)", "Potencia eléctrica (kW)"])
    st.dataframe(dfc, hide_index=True, width="stretch")
    H_n = res_new["H0"] - z_low + h_col
    H_o = res_old["H0"] - z_low + h_col
    st.write(f"Con tubería envejecida ({age_anios} años) la CDT máxima pasa de **{H_n:.2f} m** a **{H_o:.2f} m** "
             f"(+{H_o - H_n:.2f} m, {100 * (H_o - H_n) / H_n:.1f} %). Condición de dimensionamiento actual: **{cond_dis}**.")
    SHEETS["Bomba_nueva_vs_envejecida"] = dfc
    TABLAS_DOC.append(("Bomba: tubería nueva vs. envejecida", dfc[["Tubería", "Nivel de la fuente", "CDT (m)",
                       "Potencia al freno BHP (kW)", "Motor comercial (HP)"]]))
    add("Bomba (selección)", "Condición de dimensionamiento", cond_dis)
    add("Bomba (selección)", "Años de servicio / α / reducción de C", f"{age_anios} / {age_alfa} mm/año / {age_dc} %")
    add("Bomba (selección)", "CDT máx. tubería nueva", H_n, "m")
    add("Bomba (selección)", "CDT máx. tubería envejecida", H_o, "m")

    # ---- 2. potencia al freno
    st.markdown("##### 2. Potencia al freno (BHP) y datos eléctricos")
    st.caption("Potencia al freno = potencia al eje de la bomba = ρ·g·Q·CDT / η_bomba (condición de diseño).")
    e1, e2, _e3 = st.columns(3)
    volt = e1.number_input("Tensión de alimentación (V, trifásica)", 100, 13800, 440, 10)
    fp = e2.number_input("Factor de potencia del motor", 0.5, 1.0, 0.85, 0.01)
    bhp_kw = pot["p_eje"]
    hp_com_kw = pot["hp_com"] * 0.7457
    I_a = pot["p_elec"] * 1000 / (math.sqrt(3) * volt * fp)
    kwh_m3 = pot["p_elec"] / (Qlps * 3.6)
    dfp = pd.DataFrame([
        ["Caudal de diseño Qb", f"{Qlps:.2f} l/s = {Qlps * 3.6:.1f} m³/h", ""],
        ["CDT de diseño (máxima)", round(TDH_max, 2), "m"],
        ["Potencia hidráulica", round(pot["ph"], 2), "kW"],
        ["Potencia al freno BHP", f"{bhp_kw:.2f} kW = {bhp_kw / 0.7457:.1f} HP", ""],
        ["Motor requerido (con factor de servicio)", round(pot["hp_req"], 1), "HP"],
        ["Motor comercial", f"{pot['hp_com']} HP = {hp_com_kw:.1f} kW", ""],
        ["Eficiencia global bomba × motor", round(eta_b * eta_m, 3), ""],
        ["Potencia eléctrica demandada", round(pot["p_elec"], 2), "kW"],
        ["Corriente estimada", round(I_a, 1), "A"],
        ["Consumo específico de energía", round(kwh_m3, 3), "kWh/m³"],
    ], columns=["Concepto", "Valor", "Unidad"])
    st.dataframe(dfp.astype({"Valor": str}), hide_index=True, width="stretch")
    for _, r_ in dfp.iterrows():
        add("Bomba (selección)", r_["Concepto"], r_["Valor"], r_["Unidad"])

    # ---- 3. NPSH
    st.markdown("##### 3. Carga neta de succión positiva (NPSH)")
    if equipo.startswith("Centrífuga"):
        st.info("Aplica y es crítica: NPSH disponible ≥ NPSH requerido + margen (centrífuga en cárcamo seco o sobre el nivel del agua).")
    else:
        st.info("Bomba sumergida / turbina vertical: el impulsor trabaja ahogado y el NPSH rara vez es crítico; "
                "se controla con la sumergencia mínima, pero conviene verificar el NPSHr del primer impulsor.")
    z_agua_min = FUENTE_D.get("z_nda", FUENTE_D.get("namino", z_low))
    n1, n2, n3 = st.columns(3)
    temp = n1.number_input("Temperatura del agua (°C)", 0.0, 60.0, 20.0, 1.0)
    alt = n2.number_input("Altitud del sitio (msnm)", -100.0, 5000.0, float(round(z_ground0)), 10.0)
    z_eje = n3.number_input("Cota del eje / primer impulsor (msnm)", -1000.0, 6000.0,
                            float(round(FUENTE_D["z_succ"], 2)), 0.1)
    n4, n5, n6 = st.columns(3)
    hfs = n4.number_input("Pérdidas en la succión hfs (m)", 0.0, 10.0,
                          0.3 if fuente == "Pozo profundo" else float(min(h_col, 10.0)), 0.1)
    npshr = n5.number_input("NPSHr del catálogo a Qb (m); 0 si aún no lo tienes", 0.0, 100.0, 0.0, 0.1)
    marg = n6.number_input("Margen de seguridad (m) [referencia]", 0.0, 5.0, 0.6, 0.1)
    h_atm = 10.33 * (1 - 2.25577e-5 * alt) ** 5.25588
    h_vap = float(np.interp(temp, HV_T, HV_M))
    hs = z_eje - z_agua_min                  # + bomba sobre el agua / − bomba ahogada
    npshd = h_atm - h_vap - hs - hfs
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Presión atmosférica (m.c.a.)", f"{h_atm:.2f}")
    m2.metric("Presión de vapor (m.c.a.)", f"{h_vap:.2f}")
    m3.metric("Altura estática de succión hs (m)", f"{hs:+.2f}")
    m4.metric("NPSH disponible (m)", f"{npshd:.2f}")
    st.caption(f"NPSHd = Hatm − Hv − hs − hfs. Nivel mínimo del agua usado: {z_agua_min:.2f} msnm; "
               "hs > 0 si el eje queda sobre el agua, hs < 0 si la bomba está ahogada.")
    veredicto = "Captura el NPSHr del catálogo para verificar"
    if npshr > 0:
        req = npshr + marg
        hs_max = h_atm - h_vap - hfs - req
        if npshd >= req:
            st.success(f"CUMPLE: NPSHd {npshd:.2f} m ≥ NPSHr + margen = {req:.2f} m. "
                       f"Altura estática de succión máxima admisible ≈ {hs_max:+.2f} m.")
            veredicto = "Cumple"
        else:
            st.error(f"NO CUMPLE: NPSHd {npshd:.2f} m < NPSHr + margen = {req:.2f} m. Baja el eje de la bomba, "
                     f"reduce pérdidas en la succión o elige otra bomba (hs máx. admisible ≈ {hs_max:+.2f} m).")
            ALERTAS.append(f"NPSH disponible ({npshd:.2f} m) menor que NPSHr + margen ({req:.2f} m).")
            veredicto = "No cumple"
    for k_, v_, u_ in [("Presión atmosférica (altitud)", h_atm, "m.c.a."), ("Presión de vapor", h_vap, "m.c.a."),
                       ("Nivel mínimo del agua en la succión", z_agua_min, "msnm"), ("Cota eje / primer impulsor", z_eje, "msnm"),
                       ("Altura estática de succión hs", hs, "m"), ("Pérdidas en succión hfs", hfs, "m"),
                       ("NPSH disponible", npshd, "m"), ("NPSH requerido (catálogo)", npshr, "m"),
                       ("Verificación NPSH", veredicto, "")]:
        add("Bomba (NPSH)", k_, v_, u_)

    # ---- 4. curva del sistema
    st.markdown("##### 4. Curva del sistema (para superponer en la curva de la bomba)")
    qs = np.linspace(0.4, 1.4, 11) * Qlps

    def _hcol(q_m3s):
        if fuente == "Pozo profundo":
            return hf_pipe(q_m3s, d_col / 1000, prof_bomba, 0.15, 120, metodo) + 0.5
        return h_col * (q_m3s / Q) ** 2

    d_niv = z_high - z_low
    curv = {"Q (l/s)": np.round(qs, 2), "Q (m³/h)": np.round(qs * 3.6, 1)}
    for etq, e_, c_ in [("nueva", eps, Chw), ("envejecida", eps_env, C_env)]:
        h_min = np.array([hidraulica(prof, cover, D, q_ / 1000, e_, c_, metodo, acc, z_dest)["H0"] - z_low
                          + _hcol(q_ / 1000) for q_ in qs])
        curv[f"CDT {etq}, nivel mín. (m)"] = np.round(h_min, 2)
        curv[f"CDT {etq}, nivel máx. (m)"] = np.round(h_min - d_niv, 2)
    dcv = pd.DataFrame(curv)
    fcv, axc = plt.subplots(figsize=(8, 4))
    axc.plot(qs, dcv.iloc[:, 2], color="navy", label="Nueva, nivel mín.")
    axc.plot(qs, dcv.iloc[:, 3], color="navy", ls="--", label="Nueva, nivel máx.")
    axc.plot(qs, dcv.iloc[:, 4], color="red", label="Envejecida, nivel mín.")
    axc.plot(qs, dcv.iloc[:, 5], color="red", ls="--", label="Envejecida, nivel máx.")
    axc.scatter([Qlps], [H_n], color="k", zorder=5, label=f"Punto de diseño ({Qlps:.1f} l/s, {H_n:.1f} m)")
    axc.axvline(Qlps, color="gray", ls=":", lw=0.8)
    axc.set_xlabel("Caudal (l/s)")
    axc.set_ylabel("Carga dinámica total (m)")
    axc.grid(alpha=0.3)
    axc.legend(fontsize=7)
    axc.set_title("Curvas del sistema")
    fcv.tight_layout()
    q1, q2 = st.columns([1.2, 1])
    q1.pyplot(fcv)
    q2.dataframe(dcv, hide_index=True, width="stretch")
    FIGS.append(("Figura 4. Curvas del sistema (tubería nueva y envejecida; niveles mínimo y máximo de la fuente).", fcv))
    SHEETS["Curva_sistema"] = dcv

# ------------------------------------------------------------------ 8 EXPORTAR
with tabs[7]:
    resumen = pd.DataFrame(RES, columns=["Sección", "Concepto", "Valor", "Unidad"])
    resumen["Valor"] = resumen["Valor"].astype(str)
    st.dataframe(resumen, height=400, width="stretch")
    if ALERTAS:
        st.warning("**Alertas:**\n\n" + "\n".join(f"- {a}" for a in dict.fromkeys(ALERTAS)))
    SHEETS = {"Resumen": resumen, **SHEETS}
    SHEETS["Perfil_entrada"] = prof.rename(columns={"ch": "Cadenamiento (m)", "z": "Cota terreno (msnm)",
                                                 "zp": "Cota tubería (msnm)"})
    SHEETS["Accesorios_tabla"] = acc
    c1, c2, c3 = st.columns(3)
    c1.download_button("Descargar Excel", build_excel(SHEETS), "memoria_bombeo.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    c2.download_button("Descargar CSV resumen", resumen.to_csv(index=False).encode("utf-8-sig"),
                       "resumen_bombeo.csv")
    try:
        docx_bytes = build_docx(resumen, list(dict.fromkeys(ALERTAS)), FIGS,
                                f"Memoria técnica descriptiva - {proyecto}",
                                tablas=TABLAS_DOC, parrafos=PARRAFOS_DOC)
        c3.download_button("Descargar memoria Word", docx_bytes, "memoria_bombeo.docx",
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except ImportError:
        c3.info("Instala python-docx para exportar Word.")
